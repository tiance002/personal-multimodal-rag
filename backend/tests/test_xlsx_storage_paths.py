"""Offline synthetic XLSX storage regression; no ingestion/job/DB/model calls."""
import hashlib
from io import BytesIO
from pathlib import Path
import zipfile

import openpyxl
import pytest
from backend.app.adapters.parsers import ParserRegistry
from backend.app.adapters.parsers.xlsx import MEDIA, XlsxParser
from backend.app.domain.parsers import ParserError
from backend.app.domain.chunking import chunk_document
from backend.app.domain.models import ChunkRecord
from backend.app.application.citations import CitationService, InMemoryCitationStore

BASE=Path(r'E:\codex_workspace\2026-10-01\task-2\m2-table\frozen-sales.xlsx')
ACTUAL=Path(r'D:\RAG-INTEGRATION-EXEC-01-rev1attempt1\storage\objects\b1\b1587b0b8396450434aa3001d759978b8e162a07718cee815157141da813b461')

@pytest.fixture
def hashpath(tmp_path):
    raw=BASE.read_bytes();digest=hashlib.sha256(raw).hexdigest()
    path=tmp_path/'objects'/digest[:2]/digest
    path.parent.mkdir(parents=True);path.write_bytes(raw)
    assert not path.suffix
    return path

@pytest.mark.parametrize('actual',[False,True],ids=['synthetic-hashpath','retained-real-hashpath'])
def test_registry_hashpath_table_chunk_citation_matches_xlsx_and_bytes(hashpath,actual):
    path=ACTUAL if actual else hashpath
    named=Path(r'D:\RAG-INTEGRATION-EXEC-01-rev1attempt1\fixtures\DOC-INVENTORY-NORTH__v1.xlsx') if actual else BASE
    raw=path.read_bytes();assert raw==named.read_bytes()
    assert hashlib.sha256(raw).hexdigest()==path.name
    registry=ParserRegistry()
    expected=registry.parse(named,MEDIA,'doc','v1')
    document=registry.parse(path,MEDIA,'doc','v1')
    assert document.model_dump(exclude={'title'})==expected.model_dump(exclude={'title'})
    # Independent mature-library byte oracle checks formula/cache values as well.
    for cached in (False,True):
        with BytesIO(raw) as stream:
            book=openpyxl.load_workbook(stream,data_only=cached,keep_links=False)
            try:
                assert book.sheetnames==[table.sheet for table in document.tables]
                for table in document.tables:
                    for cell in table.cells:
                        value=book[table.sheet][cell.coordinate].value
                        if cell.formula:
                            assert value==(cell.cached_value if cached else cell.formula)
                        else:
                            assert XlsxParser._value(value)[0]==cell.value
            finally:book.close()
    drafts=chunk_document(document);assert drafts==chunk_document(expected)
    service=CitationService(InMemoryCitationStore())
    for i,draft in enumerate(drafts):
        assert draft.chunk_type=='table'
        assert document.markdown_content[draft.start:draft.end]==draft.content
        locator=draft.source_locator.model_dump(mode='json')
        record=ChunkRecord(str(i),'kb','doc','v1',draft.content,locator)
        detail=service.freeze('run',record)
        reread=service.resolve('run',detail.citation_id)
        assert reread.locator==locator and reread.quote==draft.content and reread.version_id=='v1'
    assert path.read_bytes()==raw

@pytest.mark.parametrize('media',['application/octet-stream','text/plain','application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/vnd.ms-excel.sheet.macroEnabled.12'])
def test_wrong_declared_mime_does_not_select_xlsx(hashpath,monkeypatch,media):
    monkeypatch.setattr(openpyxl,'load_workbook',lambda *a,**k:pytest.fail('Wrong MIME selected XLSX'))
    with pytest.raises(ParserError):ParserRegistry().parse(hashpath,media,'doc','v1')

def test_corrupt_extensionless_zip_fails_before_loader(tmp_path,monkeypatch):
    raw=b'PK\x03\x04broken archive';path=tmp_path/hashlib.sha256(raw).hexdigest();path.write_bytes(raw)
    monkeypatch.setattr(openpyxl,'load_workbook',lambda *a,**k:pytest.fail('Corrupt ZIP reached loader'))
    with pytest.raises(ParserError,match='XLSX_ARCHIVE_INVALID'):ParserRegistry().parse(path,MEDIA,'doc','v1')

@pytest.mark.parametrize('kind,error',[('merge','XLSX_SHEET_LIMIT'),('text','XLSX_TEXT_LIMIT'),('macro','XLSX_ACTIVE_CONTENT_UNSUPPORTED')])
def test_extensionless_preload_guards_have_zero_loader_calls(tmp_path,monkeypatch,kind,error):
    xml=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/><mergeCells><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>'
    if kind=='text':xml=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>'+b'x'*4097+b'</t></is></c></row></sheetData></worksheet>'
    with BytesIO() as stream:
        with zipfile.ZipFile(BASE) as source,zipfile.ZipFile(stream,'w',zipfile.ZIP_STORED) as target:
            for name in source.namelist():target.writestr(name,xml if kind!='macro' and name=='xl/worksheets/sheet1.xml' else source.read(name))
            if kind=='macro':target.writestr('xl/vbaProject.bin',b'inert')
        raw=stream.getvalue()
    path=tmp_path/hashlib.sha256(raw).hexdigest();path.write_bytes(raw)
    calls=[]
    def loader(*a,**k):calls.append(k);pytest.fail('Resource rejection reached loader')
    monkeypatch.setattr(openpyxl,'load_workbook',loader)
    with pytest.raises(ParserError,match=error):ParserRegistry().parse(path,MEDIA,'doc','v1')
    assert calls==[]

@pytest.mark.parametrize('fail_second',[False,True])
def test_binary_inputs_and_workbooks_close_on_success_and_second_load_failure(hashpath,monkeypatch,fail_second):
    original=openpyxl.load_workbook;streams=[];books=[];closed=[];modes=[]
    def loader(source,**kwargs):
        assert hasattr(source,'read') and source.read(0)==b''
        assert not source.closed
        streams.append(source);modes.append(kwargs['data_only'])
        if fail_second and len(streams)==2:raise ValueError('second-load-failure')
        book=original(source,**kwargs);books.append(book);close=book.close
        def tracked_close():closed.append(book);close()
        book.close=tracked_close
        return book
    monkeypatch.setattr(openpyxl,'load_workbook',loader)
    if fail_second:
        with pytest.raises(ParserError,match='XLSX_ARCHIVE_INVALID'):XlsxParser().parse(hashpath,'doc','v1')
    else:XlsxParser().parse(hashpath,'doc','v1')
    assert modes==[False,True] and all(stream.closed for stream in streams)
    assert closed==books
