"""Real generated format decoding; OCR/VLM/DB are explicitly SIMULATED.

No model, DB, service, network or externally configured native fixtures required.
"""
from contextlib import nullcontext
from dataclasses import replace
import hashlib
import io
import json
from pathlib import Path
import sys

import fitz
import pytest
from PIL import Image
from openpyxl import Workbook
from docx import Document
from docx.oxml import OxmlElement

from backend.app.adapters.parsers import ParserRegistry, PdfParser, classify_pdf_page
from backend.app.adapters.storage import ContentAddressedStorage
from backend.app.adapters.postgres.knowledge_repository import PostgresKnowledgeRepository
from backend.app.application.caption import CaptionEnricher
from backend.app.application.final_answer_commit import FinalAnswerCommitCheck
from backend.app.application.query_router import EvidencePlan, EvidenceTarget
from backend.app.application.structured_evidence import row_facts
from backend.app.domain.chunking import chunk_document
from backend.app.domain.models import ChunkRecord, NormalizedDocument, SourceLocator
from backend.app.domain.parsers import ParserError
from backend.tests.test_caption_fact_guard import frozen
from backend.tests.test_local_caption_enricher import FakeProvider, document
from backend.tests.test_version_source_contract import SQLRecorder

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
TARGET = EvidenceTarget('松林门店', '收入', '2025-02', '元')


def png():
    stream = io.BytesIO()
    Image.new('RGB', (40, 40), 'white').save(stream, format='PNG')
    return stream.getvalue()


def pdf_file(tmp_path, pages):
    path = tmp_path / 'generated.pdf'
    with fitz.open() as pdf:
        for kind in pages:
            page = pdf.new_page(width=200, height=200)
            if kind == 'text':
                page.insert_text((10, 30), 'Native searchable evidence.')
            else:
                page.insert_image(page.rect, stream=png())
        pdf.save(path)
    return path


@pytest.mark.parametrize('ratio,length,expected', [(.5, 100, 'scanned'), (.499, 100, 'text'),
    (.1, 9, 'scanned'), (.1, 10, 'text'), (.099, 0, 'text'), (0, 0, 'text')])
def test_fixed_weknora_thresholds_and_order(ratio, length, expected):
    assert classify_pdf_page(ratio, length) == expected


@pytest.mark.parametrize('pages', [('text',), ('scan',), ('text', 'scan')])
def test_real_pdf_page_routing_with_simulated_ocr(tmp_path, monkeypatch, pages):
    calls = []
    def ocr(self, page):
        calls.append(page.number)
        return 'SIMULATED original image words', None
    monkeypatch.setattr(PdfParser, '_ocr', ocr)
    result = ParserRegistry().parse(pdf_file(tmp_path, pages), 'application/pdf', 'doc', 'v')
    assert len(calls) == pages.count('scan')
    assert ('Native searchable' in result.markdown_content) == ('text' in pages)
    scans = [a for a in result.assets if a.asset_type == 'scanned_page']
    assert len(scans) == pages.count('scan')
    for source in scans:
        derived = next(a for a in result.assets if a.derived_from_asset_id == source.asset_id)
        assert derived.asset_type == 'ocr_text' and derived.status == 'ready'
        assert source.page_no == pages.index('scan') + 1
        assert hashlib.sha256(source.source_bytes).hexdigest() == source.source_locator['sha256']
    assert result.parser_engine == 'pymupdf/weknora-page-router-v1'
    assert result.source_mapping_available


@pytest.mark.parametrize('fault', ['ocr', 'metadata', 'enumeration', 'render'])
def test_pdf_page_failures_keep_other_native_page(tmp_path, monkeypatch, fault):
    path = pdf_file(tmp_path, ('text', 'scan'))
    def fail(*args, **kwargs):
        raise RuntimeError('SIMULATED isolated failure')
    if fault == 'ocr':
        monkeypatch.setattr(PdfParser, '_ocr', fail)
    else:
        monkeypatch.setattr(fitz.Page, {'metadata':'get_image_info', 'enumeration':'get_images', 'render':'get_pixmap'}[fault], fail)
    result = ParserRegistry().parse(path, 'application/pdf', 'doc', 'v')
    assert 'Native searchable evidence.' in result.markdown_content
    assert result.parse_status == 'partial' and result.parse_warnings


def test_empty_scanned_pdf_cannot_activate(tmp_path, monkeypatch):
    from backend.app.application.ingestion import IngestionService, IngestionWorker, InMemoryIngestionRepository
    monkeypatch.setattr(PdfParser, '_ocr', lambda *a: ('', 'OCR_EMPTY'))
    repo = InMemoryIngestionRepository()
    storage = ContentAddressedStorage(tmp_path / 'cas')
    parsers = ParserRegistry()
    raw = pdf_file(tmp_path, ('scan',)).read_bytes()
    receipt = IngestionService(repo, storage, parsers).submit_upload('kb', 'scan.pdf', 'application/pdf', io.BytesIO(raw))
    job = IngestionWorker(repo, storage, parsers).process(receipt.job_id)
    assert job.status == 'failed'
    assert repo.documents[receipt.document_id].active_version_id is None


def workbook(tmp_path, rows=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sales'
    for row in rows or [('门店','统计月份','收入（元）'), ('松林门店','2025-02',1230)]:
        sheet.append(row)
    path = tmp_path / 'generated.xlsx'
    book.save(path)
    return path


def record(draft):
    return ChunkRecord('row', 'kb', 'doc', 'v', draft.content,
                       json.loads(draft.source_locator.model_dump_json()), content_sha256=draft.content_sha256)


def test_xlsx_default_header_explicit_header_and_proof_roundtrip(tmp_path):
    path = workbook(tmp_path)
    default = ParserRegistry().parse(path, XLSX, 'doc', 'v')
    assert default.tables[0].header_rows == ()
    assert 'A: 门店' in default.markdown_content
    parsed = ParserRegistry(xlsx_first_row_as_header=True).parse(path, XLSX, 'doc', 'v')
    assert parsed.tables[0].header_rows == (1,)
    assert parsed.markdown_content == 'Sheet: Sales\n门店: 松林门店,统计月份: 2025-02,收入（元）: 1230\n'
    proof = next(p for p in parsed.table_row_proofs if p.row == 2)
    assert proof.source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert len(proof.cells) == 6 and proof.header_policy == 'explicit_first_row'
    assert NormalizedDocument.model_validate_json(parsed.model_dump_json()).table_row_proofs == parsed.table_row_proofs
    chunk = record(chunk_document(parsed)[0])
    assert row_facts(chunk.content, chunk.locator, TARGET, chunk.content_sha256)
    checker = FinalAnswerCommitCheck()
    plan = EvidencePlan('SIMULATED monthly question', targets=(TARGET,))
    assert checker.check('松林门店2025-02的收入为1230元 [E1]', frozen(chunk), plan) is None
    for bad in ['其他门店2025-02的收入为1230元 [E1]', '松林门店2025-03的收入为1230元 [E1]',
                '松林门店2025-02的收入为1230万元 [E1]', '松林门店2025-02的收入为999元 [E1]']:
        assert checker.check(bad, frozen(chunk), plan) == 'UNSUPPORTED_ANSWER'
    for mutation in ['foreign_header', 'formula', 'missing_cache', 'merged']:
        locator = json.loads(json.dumps(chunk.locator))
        cell = next(c for c in locator['cells'] if c['coordinate'] == 'C2')
        if mutation == 'foreign_header': cell['column_headers'] = ['支出（元）']
        if mutation == 'formula': cell['formula'] = '=1230'
        if mutation == 'missing_cache': cell['cache_status'] = 'missing'
        if mutation == 'merged': cell['merged_anchor'] = 'C1'
        assert not row_facts(chunk.content, locator, TARGET, chunk.content_sha256)


def test_xlsx_multisheet_merge_and_missing_formula_cache(tmp_path):
    path = workbook(tmp_path, [('same','same',None), ('merged',None,'=1+2'), ('tail',2,3)])
    book = __import__('openpyxl').load_workbook(path)
    book['Sales'].merge_cells('A2:B2')
    book.create_sheet('Second').append(['other sheet'])
    book.save(path)
    result = ParserRegistry(xlsx_first_row_as_header=True).parse(path, XLSX, 'doc', 'v')
    assert [t.sheet for t in result.tables] == ['Sales', 'Second']
    assert 'same: merged,same_2: merged' in result.markdown_content
    assert 'Sheet: Second\nA: other sheet' in result.markdown_content
    cells = {c.coordinate:c for c in result.tables[0].cells}
    assert cells['B2'].value is None and cells['B2'].merged_anchor == 'A2'
    assert cells['C2'].formula == '=1+2' and cells['C2'].cached_value is None
    assert cells['C2'].cache_status == 'missing'


def test_real_docx_declared_header_is_still_strict_source_evidence(tmp_path):
    doc = Document()
    doc.add_paragraph('Original paragraph')
    table = doc.add_table(rows=2, cols=3)
    for row, data in zip(table.rows, [('门店','统计月份','收入（元）'), ('松林门店','2025-02','1230')]):
        for cell, value in zip(row.cells, data): cell.text = value
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    path = tmp_path / 'real.docx'
    doc.save(path)
    result = ParserRegistry().parse(path, 'application/octet-stream', 'doc', 'v')
    assert 'Original paragraph' in result.markdown_content
    assert result.tables[0].header_detection == 'docx-declared-simple-header-v1'
    draft = next(c for c in chunk_document(result) if c.chunk_type == 'table' and 'R2C1:R2C3' in c.content)
    chunk = record(draft)
    assert row_facts(chunk.content, chunk.locator, TARGET, chunk.content_sha256, document_id='doc', version_id='v')
    assert all(p.source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest() for p in result.table_row_proofs)


def test_real_xls_and_corrupt_source_categories(tmp_path):
    path = Path(__file__).parent / 'fixtures/p2_weknora/ragged.xls'
    result = ParserRegistry().parse(path, 'application/octet-stream', 'doc', 'v')
    assert 'Sheet: Sheet1\nA: a,B: b,C: c' in result.markdown_content
    assert result.parse_status == 'partial' and result.parse_warnings == ('XLS_FORMULA_CACHE_UNVERIFIED',)
    assert result.table_row_proofs
    with_headers = ParserRegistry(xlsx_first_row_as_header=True).parse(path, '', 'doc', 'v')
    assert with_headers.tables[0].header_rows == (1,)
    assert 'a: d,b: e' in with_headers.markdown_content
    path = tmp_path / 'corrupt.xls'
    path.write_bytes(b'not an OLE workbook')
    with pytest.raises(ParserError) as error: ParserRegistry().parse(path, '', 'doc', 'v')
    assert error.value.category == 'SOURCE_CORRUPT'


@pytest.mark.parametrize('suffix,text', [('.txt','original\r\ntext'), ('.md','# Heading\noriginal text'),
                                        ('.html','<h1>Heading</h1><p>original text</p><script>forbidden</script>')])
def test_simple_formats_have_engine_and_source_mapping(tmp_path, suffix, text):
    path = tmp_path / ('source' + suffix)
    path.write_bytes(text.encode())
    result = ParserRegistry().parse(path, '', 'doc', 'v')
    assert 'original' in result.markdown_content and 'forbidden' not in result.markdown_content
    assert result.parser_engine != 'legacy' and result.parser_version and result.source_mapping_available
    if suffix == '.md': assert result.sections[0].heading == 'Heading'


class Guard:
    def __init__(self, allowed=True): self.permitted=allowed; self.calls=[]
    def allowed(self, doc, version): return self.permitted
    def reserve(self, doc, version): self.calls.append(('reserve', doc, version)); return 'SIMULATED-ticket'
    def settle(self, ticket, usage): self.calls.append(('settle', ticket, usage))


@pytest.mark.parametrize('permission', ['absent', 'denied', 'allowed', 'unknown_usage'])
def test_generic_cloud_caption_requires_permission_budget_and_settlement(permission):
    provider = FakeProvider()
    provider.provider_kind = 'cloud'
    guard = None if permission == 'absent' else Guard(permission != 'denied')
    if permission == 'unknown_usage': provider.error = TimeoutError()
    source = document()
    result = CaptionEnricher(provider, usage_guard=guard, enabled=True).enrich(source)
    caption = next(a for a in result.assets if a.asset_type == 'caption')
    assert result.markdown_content.startswith(source.markdown_content)
    if permission in {'absent', 'denied'}:
        assert not provider.calls and not provider.preflights and caption.status == 'failed'
        assert caption.error_code == ('VLM_UNAVAILABLE' if guard is None else 'VLM_EGRESS_DENIED')
        assert guard is None or not guard.calls
    else:
        assert guard.calls[0] == ('reserve','doc','v')
        assert guard.calls[-1] == ('settle','SIMULATED-ticket', None if permission == 'unknown_usage' else {'input_tokens':11,'output_tokens':8})
        assert caption.status == ('failed' if permission == 'unknown_usage' else 'ready')
        assert caption.derived_from_asset_id == source.assets[0].asset_id


def test_guard_backend_failure_preserves_original_and_never_sends():
    provider = FakeProvider()
    provider.provider_kind = 'cloud'
    guard = Guard()
    def unavailable(*a): raise RuntimeError('SIMULATED backend unavailable')
    guard.allowed = unavailable
    original = document()
    result = CaptionEnricher(provider, usage_guard=guard, enabled=True).enrich(original)
    assert result.markdown_content == original.markdown_content
    assert not provider.calls and not provider.preflights and not guard.calls
    assert result.assets[-1].error_code == 'VLM_UNAVAILABLE'


@pytest.mark.parametrize('ocr_failure', [False, True])
def test_image_ocr_and_caption_are_independent_and_traceable(tmp_path, monkeypatch, ocr_failure):
    path = tmp_path / 'image.png'
    path.write_bytes(png())
    monkeypatch.setattr(PdfParser, '_ocr', lambda *a: ('','OCR_UNAVAILABLE') if ocr_failure else ('SIMULATED original words', None))
    provider = FakeProvider()
    result = ParserRegistry(caption_enricher=CaptionEnricher(provider, enabled=True)).parse(path, 'image/png', 'doc', 'v')
    source, ocr, caption = result.assets
    assert ocr.derived_from_asset_id == source.asset_id == caption.derived_from_asset_id
    assert ocr.status == ('failed' if ocr_failure else 'ready') and caption.status == 'ready'
    assert bool(result.markdown_content.strip()) and source.source_bytes == path.read_bytes()
    assert [c.chunk_type for c in chunk_document(result)].count('image_caption') == 1


def test_old_locator_json_defaults_still_readable():
    old = {'kind':'table','sheet':'Sales','cell_range':'A2:C2','table_id':'old','cells':[]}
    assert SourceLocator.model_validate_json(json.dumps(old)).table_id == 'old'
    from backend.tests.test_structured_evidence import record as legacy_record
    old_chunk = legacy_record()
    assert row_facts(old_chunk.content, old_chunk.locator, TARGET, old_chunk.content_sha256)


def test_native_xls_merge_expands_only_retrieval_text():
    from backend.app.domain.models import DocumentTable, TableCell
    from backend.app.domain.table_evidence import table_row_texts
    cell = TableCell(coordinate='R1C1',row=1,column=1,value='merged',value_type='text',display='merged',row_span=2,column_span=2)
    table = DocumentTable(table_id='xls',sheet='Sheet',cell_range='R1C1:R2C2',source_format='xls',
                          row_representation='key_value',cells=[cell],start=0,end=0)
    assert table_row_texts(table) == [(1,'Sheet: Sheet\nA: merged,B: merged\n'),(2,'Sheet: Sheet\nA: merged,B: merged\n')]
    assert table.cells == [cell] and cell.row_span == 2 and cell.column_span == 2


def test_doc_converter_adapter_preserves_original_and_converted_identity(tmp_path):
    from backend.tests.test_legacy_doc import ole
    from backend.app.adapters.parsers.doc import ConvertedDocx
    doc = Document()
    doc.add_table(rows=2,cols=1).cell(0,0).text = 'converted original table'
    stream = io.BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    class Converter:
        temp_root = tmp_path
        def convert(self, raw):
            assert raw == ole()
            return ConvertedDocx(data, 'SIMULATED-converter/1')
    source = tmp_path / 'source.doc'
    source.write_bytes(ole())
    parsed = ParserRegistry(doc_converter=Converter()).parse(source, '', 'doc', 'v')
    assert parsed.parse_status == 'partial' and 'DOC_CONVERTED_LAYOUT_UNVERIFIED' in parsed.parse_warnings
    proof = parsed.table_row_proofs[0]
    assert proof.source_sha256 == hashlib.sha256(ole()).hexdigest()
    assert proof.conversion_lineage['converted_sha256'] == hashlib.sha256(data).hexdigest()
    assert proof.conversion_lineage['coordinate_basis'] == 'converted-docx-not-original-pagination'
    assert source.read_bytes() == ole()


def test_production_container_explicit_caption_injection_and_header_config(tmp_path):
    from backend.app.bootstrap import build_container
    from backend.app.config import Settings
    provider = FakeProvider()
    settings = Settings(database_url='sqlite+pysqlite:///:memory:',storage_root=tmp_path,xlsx_first_row_as_header=True)
    container = build_container(settings,model=None,caption_provider=provider)
    assert container.store.parsers.caption_enricher.provider is provider
    assert container.store.parsers.parsers[XLSX].first_row_as_header is True
    with pytest.raises(ValueError): Settings(xlsx_first_row_as_header='true')


def test_manifest_saved_with_proof_and_version_activation_sql(tmp_path, monkeypatch):
    storage = ContentAddressedStorage(tmp_path / 'cas')
    path = workbook(tmp_path)
    stored = storage.put_stream(io.BytesIO(path.read_bytes()))
    version = dict(id='v',document_id='doc',knowledge_base_id='kb',version_no=1,file_name='source.xlsx',
                   media_type=XLSX,storage_key=stored.storage_key,source_sha256=stored.sha256,original_size=stored.size)
    def handler(sql, params):
        if sql.startswith('SELECT * FROM ingestion_jobs'): return {'version_id':'v'}
        if sql.startswith('SELECT dv.*'): return version
        if sql.startswith('SELECT 1 FROM ingestion_jobs'): return (1,)
        if sql.startswith('SELECT active_version_id'): return None
    engine = SQLRecorder(handler)
    repo = PostgresKnowledgeRepository(engine, storage, parsers=ParserRegistry(xlsx_first_row_as_header=True))
    monkeypatch.setattr(repo, 'renew_job', lambda *a, **kw: True)
    monkeypatch.setattr(repo, 'update_job_progress', lambda *a, **kw: None)
    monkeypatch.setattr(repo, '_lease_heartbeat', lambda *a, **kw: nullcontext())
    monkeypatch.setattr(repo, 'get_job', lambda *a: {})
    repo.process_job('job', worker_id='worker', claim_token='claim')
    sql, params = next((sql,p) for sql,p in engine.calls if "index_status='ready'" in sql)
    assert 'processing_manifest_key=:manifest_key' in sql
    data = storage.path_for(params['manifest_key']).read_bytes()
    assert hashlib.sha256(data).hexdigest() == params['manifest_sha']
    manifest = json.loads(data)
    assert manifest['parser_engine'] == 'openpyxl' and manifest['source_mapping_available']
    assert manifest['table_row_proofs'][1]['source_sha256'] == stored.sha256
    assert any(sql.startswith('UPDATE documents SET active_version_id') for sql,p in engine.calls)
