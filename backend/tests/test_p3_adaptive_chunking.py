"""P3 offline contracts. SQL/embedding doubles are SIMULATED, PDFs decode for real."""
from contextlib import nullcontext
from dataclasses import replace
import hashlib
import importlib.util
import io
import json
from pathlib import Path
from types import SimpleNamespace

from alembic.migration import MigrationContext
from alembic.operations import Operations
import pytest

from backend.app.adapters.parsers import ParserRegistry
from backend.app.adapters.postgres.knowledge_repository import PostgresKnowledgeRepository
from backend.app.adapters.storage import ContentAddressedStorage
from backend.app.application.ingestion import InMemoryIngestionRepository, IngestionService, IngestionWorker
from backend.app.config import Settings
from backend.app.domain import adaptive_chunking as adaptive
from backend.app.domain.models import NormalizedDocument
from backend.app.domain.parsers import ParserError
from backend.app.domain.scope import Scope
from backend.app.domain.text_normalization import normalize_query
from backend.tests.test_pdf_table_evidence import chain
from backend.tests.test_version_source_contract import SQLRecorder, Result


def document(text, media_type="text/markdown"):
    return NormalizedDocument(document_id="doc", version_id="v1", title="P3",
        media_type=media_type, markdown_content=text,
        content_sha256=hashlib.sha256(text.encode()).hexdigest(), parser_version="SIMULATED/v1")


def assert_source(prepared, source):
    for chunk in prepared.parents + prepared.children:
        assert source[chunk.start:chunk.end] == chunk.content == chunk.source_locator.quote
        assert (chunk.source_locator.start, chunk.source_locator.end) == (chunk.start, chunk.end)
        assert hashlib.sha256(chunk.content.encode()).hexdigest() == chunk.content_sha256
    for child in prepared.children:
        assert len(child.content) <= 384
        expected = (child.context_header + "\n\n" if child.context_header else "") + child.content.strip()
        assert child.embedding_content == expected
        if child.parent_index is not None:
            parent = prepared.parents[child.parent_index]
            assert parent.start <= child.start < child.end <= parent.end
    for parent in prepared.parents:
        assert len(parent.content) <= 4096
    # Every character remains covered, even where semantic overlap is absent.
    covered = set()
    for child in prepared.children:
        covered.update(range(child.start, child.end))
    assert covered == set(range(len(source)))


def test_fixed_defaults_and_effective_child_overlap():
    config = adaptive.ChunkingConfig()
    assert (config.general_size, config.general_overlap, config.parent_size, config.child_size) == (512, 80, 4096, 384)
    assert config.child_size // 5 == 76
    assert (Settings().max_chunk_chars, Settings().chunk_overlap) == (512, 80)


@pytest.mark.parametrize("text,tier", [
    ("".join(f"# Heading {i}\n" + "Source statement. " * 30 + "\n" for i in range(4)), "heading"),
    ("第一章 数据\n" + "资料数值为 -7.25。" * 65 + "\n第二节 证据\n" + "应核对原文。" * 65, "heuristic"),
    ("OCR line one.\nOCR line two.\n\fPage 2\n" + "Scanned content. " * 100, "heuristic"),
    ("连续中文文本应保留句末边界。" * 100, "legacy"),
    ("UNBROKEN" * 180, "legacy"),
])
def test_adaptive_profiles_exact_source_and_fallback(text, tier):
    prepared = adaptive.prepare_document(document(text))
    assert prepared.diagnostics[0]["selected"] == tier
    assert_source(prepared, text)


def test_marker_thresholds_and_fenced_false_headings():
    text = "```md\n# fake\n# fake2\n# fake3\n第一章 假标题\n```\nordinary"
    profile = adaptive.profile_text(text)
    assert profile["headings"] == profile["chapters"] == 0
    assert adaptive.select_tiers(profile, "auto") == ["legacy"]
    profile = adaptive.profile_text("\n".join(f"{i}. OCR body" for i in range(1, 6)))
    assert adaptive.select_tiers(profile, "auto") == ["heuristic", "legacy"]


def test_unterminated_ocr_fence_and_cross_line_hash_are_not_headings():
    source = "#\nordinary text\n```md\n# fake\n# fake2\n# fake3\n第一章 假标题"
    profile = adaptive.profile_text(source)
    assert profile["headings"] == profile["chapters"] == 0
    prepared = adaptive.prepare_document(document(source))
    assert all(not c.context_header for c in prepared.children)


def test_heading_coalescing_uses_common_breadcrumb_not_a_sibling_title():
    source = "# Root\n## First\n" + "one " * 15 + "\n## Second\n" + "two " * 15 + "\n## Third\n" + "three " * 15
    prepared = adaptive.prepare_document(document(source))
    first = prepared.children[0]
    assert "## First" in first.content and "## Second" in first.content
    assert first.context_header == "# Root"
    assert_source(prepared,source)


def test_very_long_heading_is_bounded_without_fabricating_or_truncating_quote():
    source = "# " + "标题" * 140 + "\n" + "源文。" * 600
    prepared = adaptive.prepare_document(document(source))
    assert_source(prepared,source)
    source = "# " + "标题" * 300 + "\nsource"
    prepared = adaptive.prepare_document(document(source))
    assert_source(prepared, source)
    assert all(c.context_header == source.splitlines()[0] for c in prepared.children)
    assert all(len(c.embedding_content) > 512 for c in prepared.children)


def test_rejected_tier_records_reason_and_real_recursive_fallback(monkeypatch):
    original = adaptive._tier_ranges
    def failing_tier(text, size, overlap, tier, profile):
        return [] if tier == "heading" else original(text, size, overlap, tier, profile)
    monkeypatch.setattr(adaptive, "_tier_ranges", failing_tier)  # SIMULATED tier failure only.
    source = "# A\na\n# B\nb\n# C\nc\n" + "body " * 100
    ranges, diagnostics = adaptive.split_text(source, 384, 76)
    assert diagnostics["rejected"] == [{"tier": "heading", "reason": "no chunks produced"}]
    assert diagnostics["selected"] == "legacy" and ranges
    assert adaptive.validate_ranges(source, ranges, 384) is None


def test_empty_final_output_is_an_explicit_failure(monkeypatch):
    monkeypatch.setattr(adaptive, "_tier_ranges", lambda *args: [])
    with pytest.raises(ParserError, match="CHUNK_VALIDATION_FAILED"):
        adaptive.split_text("nonempty source", 384, 76)
    with pytest.raises(ParserError, match="EMPTY_TEXT"):
        adaptive.prepare_document(document(" \n"))


def test_blank_page_separators_do_not_generate_empty_embedding_inputs():
    source = "\n" * 900 + "actual source -7.25"
    prepared = adaptive.prepare_document(document(source,"text/plain"))
    assert prepared.children and all(c.content.strip() and c.embedding_content.strip() for c in prepared.children)
    assert any(d.get("skipped_whitespace",0)>0 for d in prepared.diagnostics)
    assert any("actual source -7.25" in c.source_locator.quote for c in prepared.children)
    assert all(source[c.start:c.end] == c.content == c.source_locator.quote for c in prepared.children)


@pytest.mark.parametrize("protected", [
    "| A | B |\n| --- | --- |\n| -7.25 | 123 |\n",
    "```python\n# hidden heading\nprint(-7.25)\n```",
    "$$x = -7.25 + y$$", "[原文链接](https://example.test/source)",
    "![图像](https://example.test/image.png)", "`inline(-7.25)`",
])
def test_small_protected_span_remains_whole(protected):
    source = "前置正文。" * 60 + "\n" + protected + "\n" + "后续正文。" * 100
    prepared = adaptive.prepare_document(document(source))
    start, end = source.index(protected), source.index(protected) + len(protected)
    assert any(c.start <= start and c.end >= end for c in prepared.children)
    assert all(not start < edge < end for c in prepared.children for edge in (c.start, c.end))
    assert_source(prepared, source)


@pytest.mark.parametrize("source", ["```\n" + "code" * 600 + "\n```", "$$" + "数" * 2000 + "$$"])
def test_oversized_protection_falls_back_to_bounded_source_spans(source):
    prepared = adaptive.prepare_document(document(source))
    assert len(prepared.children) > 1
    assert_source(prepared, source)


def test_breadcrumb_embedding_and_citation_are_separate():
    source = "# Root\n" + "root source. " * 90 + "\n## Child\n" + "exact -7.25 source. " * 100
    prepared = adaptive.prepare_document(document(source))
    nested = next(c for c in prepared.children if c.start > source.index("## Child") and c.context_header.endswith("## Child"))
    assert nested.context_header == "# Root\n## Child"
    assert nested.embedding_content == nested.context_header + "\n\n" + nested.content.strip()
    assert "# Root" not in nested.source_locator.quote
    assert_source(prepared, source)


def test_parent_children_and_short_equivalent_parent_omitted():
    short = adaptive.prepare_document(document("# Short\nsource -7.25"))
    assert short.parents == [] and short.children[0].parent_index is None
    source = "# Long\n" + "source sentence -7.25。" * 600
    long = adaptive.prepare_document(document(source))
    assert len(long.parents) >= 3 and all(c.parent_index is not None for c in long.children)
    assert_source(long, source)


def test_preview_and_memory_ingestion_use_same_parser_and_configuration(tmp_path):
    storage = ContentAddressedStorage(tmp_path)
    parsers = ParserRegistry()
    repo = InMemoryIngestionRepository()
    service = IngestionService(repo, storage, parsers)
    source = ("# Scope\n" + "source -7.25. " * 200).encode()
    receipt = service.submit_upload("kb", "same.md", "text/markdown", io.BytesIO(source))
    assert IngestionWorker(repo, storage, parsers).process(receipt.job_id).status == "succeeded"
    version = repo.versions[receipt.version_id]
    preview = adaptive.preview_document_chunks(version.normalized_document)
    assert preview.children == version.chunks and preview.parents == version.parents
    assert preview.config.identity == version.index_identity


@pytest.mark.parametrize("change", [dict(child_size=383), dict(parent_size=4095), dict(general_overlap=79), dict(strategy="legacy")])
def test_profile_identity_includes_input_semantics(change):
    base = adaptive.ChunkingConfig()
    changed = replace(base, **change)
    assert base.identity != changed.identity
    assert adaptive.embedding_fingerprint("bge-m3:latest", 1024, base) != adaptive.embedding_fingerprint("bge-m3:latest", 1024, changed)
    assert adaptive.embedding_fingerprint("bge-m3:latest", 1024, base) != "ollama-local-1024"


def test_schema_and_chunker_versions_are_part_of_profile_identity(monkeypatch):
    config = adaptive.ChunkingConfig()
    original = config.identity
    monkeypatch.setattr(adaptive,"SCHEMA_VERSION","SIMULATED/future-schema")
    assert config.identity != original
    changed_schema = config.identity
    monkeypatch.setattr(adaptive,"CHUNKER_VERSION","SIMULATED/future-chunker")
    assert config.identity != changed_schema


def test_native_row_above_child_target_keeps_complete_cell_evidence(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "atomic.xlsx"
    book = Workbook()
    book.active.append(["Name", "Value"])
    book.active.append(["source", "original " * 100])
    book.save(path)
    parsed = ParserRegistry().parse(path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "doc", "v1")
    assert parsed.tables and parsed.tables[0].cells
    before = parsed.model_dump()
    from backend.app.domain.chunking import chunk_document
    expected = chunk_document(parsed, max_chars=4096, overlap=0)
    prepared = adaptive.prepare_document(parsed)
    assert len(prepared.children) == len(expected)
    for actual, original in zip(prepared.children, expected, strict=True):
        assert actual.content == original.content
        assert actual.content_sha256 == original.content_sha256
        assert actual.source_locator == original.source_locator
        assert actual.parent_index is None
    assert any(len(c.embedding_content) > 512 for c in prepared.children)
    assert parsed.model_dump() == before


def test_real_frozen_pdf_rows_keep_geometry_numbers_and_citations(tmp_path):
    path = Path(__file__).parent / "fixtures/p2_pdf_closure/tables.pdf"
    before = path.read_bytes()
    parsed = ParserRegistry(tessdata=tmp_path / "absent").parse(path, "application/pdf", "doc", "v1")
    prepared = adaptive.prepare_document(parsed)
    tables = [c for c in prepared.children if c.chunk_type == "table"]
    assert tables
    for row in tables:
        assert row.parent_index is None and row.source_locator.cells
        assert row.content == parsed.markdown_content[row.start:row.end] == row.source_locator.quote
        _, citation = chain(parsed, row, row.content, f"p3-{row.chunk_index}")
        assert citation["page"] == row.source_locator.page
        assert citation["bbox"] == list(row.source_locator.bbox)
        assert citation["quote"] == row.content
        assert row.source_locator.table_id in {t.table_id for t in parsed.tables}
    assert path.read_bytes() == before


class Embeddings:
    """SIMULATED model; no network or real inference."""
    embedding_model = "bge-m3:latest"
    def __init__(self): self.inputs = []
    def embed(self, texts, **kwargs):
        self.inputs.extend(texts)
        return SimpleNamespace(vectors=[[float(i)] * 1024 for i in range(len(texts))],
                               model=self.embedding_model, dimensions=1024)


def production_job(tmp_path, monkeypatch, *, ready=False):
    storage = ContentAddressedStorage(tmp_path / "cas")
    raw = ("# Root\n" + "source amount -7.25。" * 320).encode()
    stored = storage.put_stream(io.BytesIO(raw))
    version = dict(id="v1", document_id="doc", knowledge_base_id="kb", version_no=1,
        file_name="source.md", media_type="text/markdown", storage_key=stored.storage_key,
        source_sha256=stored.sha256, original_size=stored.size, index_status="ready" if ready else "queued")
    def handler(sql, params):
        if sql.startswith("SELECT * FROM ingestion_jobs"): return {"version_id": "v1"}
        if sql.startswith("SELECT dv.*"): return version
        if sql.startswith("SELECT 1 FROM ingestion_jobs"): return (1,)
        if sql.startswith("INSERT INTO embedding_profiles"): return params["id"]
    engine = SQLRecorder(handler)
    # SIMULATED connection protocol only; independent real-connection races
    # are covered by test_clean_slate_model_profiles, not this recorder.
    engine.get_isolation_level = lambda: "READ COMMITTED"
    model = Embeddings()
    repo = PostgresKnowledgeRepository(engine, storage, embedding_provider=model)
    monkeypatch.setattr(repo, "renew_job", lambda *a, **kw: True)
    monkeypatch.setattr(repo, "update_job_progress", lambda *a, **kw: None)
    monkeypatch.setattr(repo, "_lease_heartbeat", lambda *a, **kw: nullcontext())
    monkeypatch.setattr(repo, "get_job", lambda *a: {})
    repo.process_job("job", worker_id="worker", claim_token="claim")
    return engine, model, repo, version


def test_production_persists_source_context_and_child_only_indexes(tmp_path, monkeypatch):
    engine, model, repo, version = production_job(tmp_path, monkeypatch)
    rows = [p for sql,p in engine.calls if sql.startswith("INSERT INTO chunks")]
    parents = {p["id"]: p for p in rows if p["chunk_role"] == "parent"}
    children = [p for p in rows if p["chunk_role"] == "child"]
    assert parents and children and len(model.inputs) == len(children)
    assert len({p["chunk_index"] for p in rows}) == len(rows)
    assert [p["chunk_index"] for p in children] == list(range(len(children)))
    assert {p["chunk_index"] for p in parents.values()} == set(range(len(children),len(rows)))
    raw = repo.storage.path_for(version["storage_key"]).read_text(encoding="utf-8")
    for child, embedding in zip(children, model.inputs, strict=True):
        parent = parents[child["parent_id"]]
        assert child["version_id"] == parent["version_id"]
        assert parent["start_pos"] <= child["start_pos"] < child["end_pos"] <= parent["end_pos"]
        assert child["content"] == raw[child["start_pos"]:child["end_pos"]]
        assert json.loads(child["locator"])["quote"] == child["content"]
        assert embedding == child["context_header"] + "\n\n" + child["content"].strip()
    indexed = {p["chunk_id"] for sql,p in engine.calls if sql.startswith("INSERT INTO chunk_terms")}
    vectors = [p for sql,p in engine.calls if sql.startswith("INSERT INTO chunk_embeddings")]
    assert indexed == {p["id"] for p in children} == {p["chunk_id"] for p in vectors}
    assert not indexed.intersection(parents)
    profile = next(p for sql,p in engine.calls if sql.startswith("INSERT INTO embedding_profiles"))
    profile_sql = next(sql for sql,_ in engine.calls if sql.startswith("INSERT INTO embedding_profiles"))
    assert "ON CONFLICT (provider,model_name,model_revision,dimension,distance,fingerprint) DO NOTHING RETURNING id" in profile_sql
    assert profile["fingerprint"] == adaptive.embedding_fingerprint(model.embedding_model, 1024, repo.chunking_config)
    # R3 records provider revision honestly; the P3 index identity remains in
    # the fingerprint and chunk/version columns, rather than posing as revision.
    assert profile["revision"] == "UNKNOWN"
    update = next(p for sql,p in engine.calls if "index_status='ready'" in sql)
    manifest = json.loads(repo.storage.path_for(update["manifest_key"]).read_text(encoding="utf-8"))
    assert manifest["chunking_diagnostics"] and manifest["index_identity"] == update["index_identity"]
    assert any(sql.startswith("UPDATE documents SET active_version_id") for sql,_ in engine.calls)
    assert not any("DELETE FROM" in sql for sql,_ in engine.calls)


def test_indexed_version_reprocessing_cannot_overwrite_history(tmp_path, monkeypatch):
    engine, model, repo, _ = production_job(tmp_path, monkeypatch, ready=True)
    assert not model.inputs
    assert any("REINDEX_REQUIRES_NEW_VERSION" in sql for sql,_ in engine.calls)
    assert not any(sql.startswith("INSERT") or sql.startswith("UPDATE document_versions") or sql.startswith("UPDATE documents ") for sql,_ in engine.calls)


class Rows(Result):
    def __iter__(self): return iter(self.value or [])
    def scalars(self): return self


class ReadSQL(SQLRecorder):
    def execute(self, statement, params=None):
        sql = " ".join(str(statement).split())
        self.calls.append((sql, params or {}))
        return Rows(self.handler(sql, params or {}))


def test_production_read_filters_new_identity_and_legacy_profile(tmp_path):
    # SIMULATED SQL result: a legacy profile cannot satisfy the new fingerprint.
    legacy = {"fingerprint": "ollama-local-1024", "id": "old-profile"}
    engine = ReadSQL(lambda sql,p: legacy["id"] if p.get("fingerprint") == legacy["fingerprint"] else None)
    repo = PostgresKnowledgeRepository(engine, ContentAddressedStorage(tmp_path))
    assert repo.get_embedding_profile_id("bge-m3:latest", 1024) is None
    scope = Scope(("kb",))
    assert repo.keyword_candidates(scope, normalize_query("source"), 10) == []
    assert repo.vector_candidates(scope, [0.] * 1024, 10, profile_id="old-profile") == []
    assert repo.list_active_chunks(scope) == []
    for sql, params in engine.calls[1:]:
        assert "c.chunk_role='child'" in sql and "c.index_identity=:index_identity" in sql
        assert "dv.index_identity=:index_identity" in sql
        assert params["index_identity"] == repo.chunking_config.identity
        assert "active_version_id = c.version_id" in sql and "c.knowledge_base_id IN" in sql
    assert "ep.fingerprint=:fingerprint" in engine.calls[2][0]


def test_document_views_exclude_duplicate_parents_but_history_readback_stays_open(tmp_path):
    legacy = dict(id="old-child", knowledge_base_id="kb", document_id="doc", version_id="old-version",
                  content="historical -7.25", content_sha256="a"*64, locator={"quote":"historical -7.25"})
    def handler(sql,params):
        if sql.startswith("SELECT active_version_id"): return ("old-version",)
        if sql.startswith("SELECT content FROM chunks"): return [legacy["content"]]
        if "FROM chunks WHERE id=:id" in sql: return legacy
    engine = ReadSQL(handler)
    repo = PostgresKnowledgeRepository(engine,ContentAddressedStorage(tmp_path))
    assert repo.list_chunks("doc") == []
    assert repo.get_document_content("doc") == legacy["content"]
    assert repo.get_chunk("old-child").content == legacy["content"]
    assert "chunk_role='child'" in engine.calls[0][0] and "index_identity" not in engine.calls[0][0]
    assert "chunk_role='child'" in engine.calls[2][0] and "index_identity" not in engine.calls[2][0]
    assert "index_identity" not in engine.calls[3][0] and "active_version" not in engine.calls[3][0]


def test_graph_source_reader_cannot_turn_stored_parents_into_graph_recall(tmp_path):
    from backend.app.adapters.postgres.graph_repository import PostgresGraphRepository
    engine = ReadSQL(lambda sql,params: [])  # SIMULATED SQL only.
    repo = PostgresGraphRepository(engine,ContentAddressedStorage(tmp_path))
    assert repo.list_version_chunks("doc","v1") == []
    sql,params = engine.calls[0]
    assert "chunk_role='child'" in sql
    assert "document_id=:document_id AND version_id=:version_id" in sql
    assert params == {"document_id":"doc","version_id":"v1"}


def test_additive_migration_compiles_without_db_and_keeps_same_version_fk():
    path = Path("alembic/versions/0016_parent_child_chunks.py")
    spec = importlib.util.spec_from_file_location("p3_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    output = io.StringIO()
    context = MigrationContext.configure(dialect_name="postgresql", opts={"as_sql": True, "output_buffer": output})
    with Operations.context(context):
        migration.upgrade()
    sql = output.getvalue()
    assert "FOREIGN KEY(version_id, parent_id) REFERENCES chunks (version_id, id)" in sql
    assert "DEFAULT 'legacy/v1' NOT NULL" in sql
    assert "context_header TEXT DEFAULT '' NOT NULL" in sql
    assert "DROP" not in sql and "DELETE" not in sql and "UPDATE" not in sql
    with pytest.raises(RuntimeError, match="Forward-only"):
        migration.downgrade()
