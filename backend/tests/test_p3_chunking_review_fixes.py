"""P3 review counterexamples: real source slicing and real XLSX decoding, no model/DB."""
import hashlib

from openpyxl import Workbook
import pytest

from backend.app.adapters.parsers import ParserRegistry
from backend.app.domain import adaptive_chunking as adaptive
from backend.app.domain.chunking import chunk_document
from backend.tests.test_p3_adaptive_chunking import document


@pytest.mark.parametrize("strategy", ["auto", "heading", "heuristic", "legacy"])
@pytest.mark.parametrize("row_count", [32, 240])
def test_long_markdown_table_short_rows_stay_whole(strategy, row_count):
    prefix = "# Report\n# Table\n"
    header = "| Name | Value |\n| --- | --- |\n"
    rows = [f"| row-{i:04d} | -{i}.25 |\n" for i in range(row_count)]
    source = prefix + header + "".join(rows) + "# End\nUnmodified closing source."
    prepared = adaptive.prepare_document(document(source), adaptive.ChunkingConfig(strategy=strategy))
    # Complete coverage and exact offsets/quotes, including table whitespace.
    covered = set()
    for child in prepared.children:
        assert source[child.start:child.end] == child.content == child.source_locator.quote
        assert (child.start, child.end) == (child.source_locator.start, child.source_locator.end)
        assert child.content_sha256 == hashlib.sha256(child.content.encode()).hexdigest()
        covered.update(range(child.start, child.end))
    assert covered == set(range(len(source)))
    for parent in prepared.parents:
        assert parent.content == source[parent.start:parent.end] == parent.source_locator.quote
    # The header+separator is one protected unit; each body row is another.
    start = len(prefix)
    units = [header] + rows
    for unit in units:
        end = start + len(unit)
        assert len(unit) < 384
        assert any(c.start <= start and end <= c.end for c in prepared.children), (strategy, start, unit)
        assert all(not start < edge < end for c in prepared.children for edge in (c.start, c.end)), (strategy, start, unit)
        assert all(not start < edge < end for c in prepared.parents for edge in (c.start, c.end))
        start = end


def native_document(tmp_path, value_length, name="original"):
    path = tmp_path / "native-atomic.xlsx"
    book = Workbook()
    book.active.append(["Name", "Value"])
    book.active.append([name, "v" * value_length])
    book.save(path)
    return ParserRegistry().parse(path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "doc", "v1")


@pytest.mark.parametrize("general,child", [(128, 64), (512, 384)])
def test_native_proof_above_targets_is_not_an_embedding_capacity_violation(tmp_path, general, child):
    parsed = native_document(tmp_path, 900)
    before = parsed.model_dump()
    config = adaptive.ChunkingConfig(general_size=general, child_size=child)
    expected = chunk_document(parsed, max_chars=config.parent_size, overlap=0)
    prepared = adaptive.prepare_document(parsed, config)
    assert len(prepared.children) == len(expected)
    assert any(len(c.embedding_content) > 512 for c in prepared.children)
    for original, actual in zip(expected, prepared.children, strict=True):
        assert actual.content == original.content
        assert actual.content_sha256 == original.content_sha256
        assert actual.source_locator == original.source_locator
        assert actual.source_locator.cells and actual.parent_index is None
    assert parsed.model_dump() == before


def test_long_breadcrumb_is_separate_from_child_size_and_preserved_verbatim():
    header = "# " + "标题" * 300
    source = header + "\n" + "原始数值 -7.25。" * 90
    small = adaptive.prepare_document(document(source), adaptive.ChunkingConfig(general_size=128))
    default = adaptive.prepare_document(document(source))
    assert [(c.start, c.end, c.content) for c in small.children] == [(c.start, c.end, c.content) for c in default.children]
    for child in default.children:
        assert len(child.content) <= 384
        assert child.context_header == header
        assert child.embedding_content == header + "\n\n" + child.content.strip()
        assert len(child.embedding_content) > 512  # A target size is not a model input limit.
        assert child.content == source[child.start:child.end] == child.source_locator.quote
    assert adaptive.ChunkingConfig().identity != "p3:14078be5f947fc1274996fc2aefdda17e402ed1a85e1cdf0b52fceb87cfc472d"


def test_existing_native_admission_overflow_still_rejects_without_cutting_proof(tmp_path):
    # Each cell is below the XLSX parser's 4096-character field admission.
    # Together the serialized row exceeds the existing 4096 row admission.
    parsed = native_document(tmp_path, 2500, name="n" * 2500)
    before = parsed.model_dump()
    with pytest.raises(ValueError, match="TABLE_ROW_TOO_LARGE"):
        chunk_document(parsed, max_chars=4096, overlap=0)
    with pytest.raises(ValueError, match="TABLE_ROW_TOO_LARGE"):
        adaptive.prepare_document(parsed)
    assert parsed.model_dump() == before


@pytest.mark.parametrize("row_length", [4096, 4097])
def test_native_row_admission_exact_character_boundary(tmp_path, row_length):
    # This is the existing source-row admission, NOT a model token capacity.
    from backend.app.domain.table_evidence import table_row_texts
    name = "n" * 2048
    probe = native_document(tmp_path, 1, name=name)
    overhead = len(dict(table_row_texts(probe.tables[0]))[2]) - 1
    parsed = native_document(tmp_path, row_length - overhead, name=name)
    row = dict(table_row_texts(parsed.tables[0]))[2]
    assert len(row) == row_length
    before = parsed.model_dump()
    if row_length == 4097:
        with pytest.raises(ValueError, match="TABLE_ROW_TOO_LARGE"):
            adaptive.prepare_document(parsed)
    else:
        prepared = adaptive.prepare_document(parsed)
        actual = next(c for c in prepared.children if c.content == row)
        assert actual.source_locator.cells and actual.parent_index is None
        assert parsed.markdown_content[actual.start:actual.end] == actual.source_locator.quote == row
        assert actual.content_sha256 == hashlib.sha256(row.encode()).hexdigest()
    assert parsed.model_dump() == before
