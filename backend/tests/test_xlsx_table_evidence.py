import json,zipfile
from pathlib import Path
import pytest
from backend.app.adapters.parsers import XlsxParser
from backend.app.domain.parsers import ParserError
from backend.app.domain.chunking import chunk_document
from backend.app.domain.models import NormalizedDocument,ChunkRecord
from backend.app.application.retrieval import HybridRetriever,InMemoryRetrievalRepository
from backend.app.domain.scope import Scope
from backend.app.application.citations import CitationService,InMemoryCitationStore

FIXTURE=Path(r'E:\codex_workspace\2026-10-01\task-2\m2-table\frozen-sales.xlsx')

def parsed():return XlsxParser().parse(FIXTURE,'doc','v1')
def cells(document,sheet):return {c.coordinate:c for c in next(t for t in document.tables if t.sheet==sheet).cells}

def test_names_sparse_coordinates_merges_and_block_roundtrip():
 d=parsed();assert [t.sheet for t in d.tables]==['Sales','Inventory']
 assert d.tables[0].cell_range=='A1:E6' and d.tables[0].merged_ranges==('A1:E1',)
 assert d.tables[0].header_rows==() and d.tables[1].header_rows==()
 assert all(t.header_detection=='column_letters' for t in d.tables)
 assert cells(d,'Sales')['B1'].merged_anchor=='A1'
 assert cells(d,'Inventory')['B2'].value is None and cells(d,'Inventory')['C2'].value==7
 assert all(cells(d,'Sales')[f'{col}5'].value is None for col in 'ABCDE')
 roundtrip=NormalizedDocument.model_validate_json(d.model_dump_json())
 assert roundtrip.tables==d.tables and roundtrip.blocks==d.blocks

def test_typed_formats_formula_and_cache_are_separate():
 c=cells(parsed(),'Sales')
 assert c['B4'].value==-30.125 and c['B4'].raw_number=='-30.125' and c['B4'].number_format=='0.000'
 assert c['D3'].value==0.125 and c['D3'].display=='12.5%'
 assert c['E3'].value=='2026-10-01T00:00:00' and c['E3'].value_type=='datetime'
 assert c['C6'].formula=='=SUM(C3:C4)' and c['C6'].cache_status=='present' and c['C6'].cached_value==60
 assert c['D6'].formula=='=1/8' and c['D6'].cached_value is None and c['D6'].cache_status=='missing'
 assert 'UNKNOWN_CACHE_MISSING' in c['D6'].display

@pytest.mark.parametrize('sheet,coordinate,header,expected',[('Sales','C3','C',20),('Sales','B4','B',-30.125),('Sales','D3','D',0.125),('Sales','E3','E','2026-10-01T00:00:00'),('Sales','C6','C',60),('Sales','D6','D',None),('Inventory','C2','C',7)])
def test_question_evidence_header_and_citation_json_chain(sheet,coordinate,header,expected):
 d=parsed();drafts=chunk_document(d)
 row=int(''.join(ch for ch in coordinate if ch.isdigit()))
 draft=next(c for c in drafts if c.source_locator.sheet==sheet and c.source_locator.cell_range.startswith('A'+str(row)+':'))
 assert draft.chunk_type=='table'
 if expected is not None:assert header in draft.content
 else:assert 'UNKNOWN_CACHE_MISSING' not in draft.content  # no fabricated retrieval value
 assert d.markdown_content[draft.start:draft.end]==draft.content
 persisted=json.loads(json.dumps({'content':draft.content,'locator':draft.source_locator.model_dump(mode='json')}))
 target=next(c for c in persisted['locator']['cells'] if c['coordinate']==coordinate)
 assert header in target['column_headers']
 assert (target['cached_value'] if target['formula'] else target['value'])==expected
 repo=InMemoryRetrievalRepository();record=ChunkRecord('row','kb','doc','v1',persisted['content'],persisted['locator'])
 repo.add(record);repo.add(ChunkRecord('private','other','doc2','v1',record.content,record.locator))
 repo.add(ChunkRecord('old','kb','doc','v0',record.content,record.locator,is_current=False))
 result=HybridRetriever(repo).retrieve(Scope.from_ids(['kb']),header if expected is not None else 'Total')
 assert [x.chunk.chunk_id for x in result.items]==['row']
 service=CitationService(InMemoryCitationStore());detail=service.freeze('run',result.items[0].chunk)
 reread=service.resolve('run',detail.citation_id)
 assert reread.locator['sheet']==sheet and reread.locator['cell_range']==draft.source_locator.cell_range
 assert reread.quote==draft.content and reread.version_id=='v1'
 from backend.app.application.context_builder import ContextBuilder
 context,labels=ContextBuilder().build('run',result.items,service)
 assert labels==['E1'] and draft.content in context
 from fastapi.testclient import TestClient
 from backend.app.main import create_app
 from backend.app.config import Settings
 from types import SimpleNamespace
 class FakeStore:
  def get_citation(self,run,label):
   detail=service.resolve(run,label)
   return {'citation_id':detail.citation_id,'quote':detail.quote,'version_id':detail.version_id,'locator':detail.locator,'current_status':detail.current_status}
 settings=Settings()
 client=TestClient(create_app(container=SimpleNamespace(settings=settings,store=FakeStore())))
 response=client.get('/api/v1/runs/run/citations/E1')
 assert response.status_code==200
 assert response.json()['data']['locator']==persisted['locator']

def mutated(tmp_path,xml=None,extra=None):
 path=tmp_path/'bad.xlsx'
 with zipfile.ZipFile(FIXTURE) as source,zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as out:
  for name in source.namelist():out.writestr(name,xml if xml is not None and name=='xl/worksheets/sheet1.xml' else source.read(name))
  if extra:out.writestr(*extra)
 return path

def test_xml_entity_expansion_is_rejected(tmp_path):
 xml=b'<!DOCTYPE worksheet [<!ENTITY x "boom"><!ENTITY y "&x;&x;&x;">]><worksheet>&y;</worksheet>'
 with pytest.raises(ParserError,match='XLSX_UNSAFE_XML'):XlsxParser().parse(mutated(tmp_path,xml),'doc','v1')

@pytest.mark.parametrize('extra',[('xl/vbaProject.bin',b'never run'),('xl/externalLinks/externalLink1.xml',b'<x/>')])
def test_macro_or_external_workbook_rejected(tmp_path,extra):
 with pytest.raises(ParserError,match='XLSX_ACTIVE_CONTENT_UNSUPPORTED'):XlsxParser().parse(mutated(tmp_path,extra=extra),'doc','v1')

def test_zip_expansion_limit(tmp_path):
 with pytest.raises(ParserError,match='XLSX_ARCHIVE_LIMIT'):XlsxParser().parse(mutated(tmp_path,extra=('padding.bin',b'0'*1000000)),'doc','v1')

def test_oversized_sheet_rejected_before_openpyxl_load(tmp_path,monkeypatch):
 import openpyxl
 monkeypatch.setattr(openpyxl,'load_workbook',lambda *a,**k:pytest.fail('Must reject before loading'))
 xml=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:XFD1048576"/></worksheet>'
 with pytest.raises(ParserError,match='XLSX_SHEET_LIMIT'):XlsxParser().parse(mutated(tmp_path,xml),'doc','v1')

def test_oversized_row_not_silently_split_or_lost():
 with pytest.raises(ValueError,match='TABLE_ROW_TOO_LARGE'):chunk_document(parsed(),max_chars=20,overlap=0)


def test_workbook_order_does_not_invent_sheet_identity(tmp_path):
 from defusedxml.ElementTree import fromstring
 from xml.etree.ElementTree import tostring
 path=tmp_path/'reordered.xlsx'
 with zipfile.ZipFile(FIXTURE) as source,zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as out:
  for name in reversed(source.namelist()):
   data=source.read(name)
   if name=='xl/workbook.xml':
    root=fromstring(data);sheets=root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets')
    sheets[:]=list(reversed(list(sheets)));data=tostring(root)
   out.writestr(name,data)
 d=XlsxParser().parse(path,'doc','v1')
 assert [t.sheet for t in d.tables]==['Inventory','Sales']
 assert cells(d,'Sales')['B4'].raw_number=='-30.125' and cells(d,'Inventory')['C2'].value==7

def test_mixed_unhandled_blocks_do_not_silently_disappear():
 d=parsed();d=d.model_copy(update={'markdown_content':d.markdown_content+'Original text must survive'})
 with pytest.raises(ValueError,match='TABLE_BLOCK_COVERAGE_REQUIRED'):chunk_document(d)

def test_plain_legacy_locator_payload_remains_compatible():
 from backend.app.domain.models import SourceLocator
 locator=SourceLocator.model_validate({'kind':'pdf','page':2,'quote':'legacy'})
 assert locator.kind=='pdf' and locator.page==2 and locator.sheet is None


def test_repeated_header_text_has_explicit_bound(tmp_path):
 from openpyxl import Workbook
 workbook=Workbook();sheet=workbook.active;sheet['A1']='Header '+('abcd'*150);sheet['B1']='Second';sheet['A2']=1;sheet['B2']=2
 path=tmp_path/'large-header.xlsx';workbook.save(path)
 with pytest.raises(ParserError,match='XLSX_TEXT_LIMIT'):XlsxParser(first_row_as_header=True).parse(path,'doc','v1')
