from __future__ import annotations

import hashlib
import re
import posixpath
import zipfile
from pathlib import Path
from datetime import datetime,date,time

import openpyxl
from openpyxl.utils.cell import range_boundaries,get_column_letter
from openpyxl.xml import functions as xml_functions
from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError

from backend.app.domain.models import NormalizedDocument,DocumentTable,TableCell,DocumentBlock,DocumentSection,SourceLocator
from backend.app.domain.parsers import ParserError
from backend.app.domain.table_evidence import table_row_texts

MEDIA='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NS='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

class XlsxParser:
    """Bounded, inert workbook reader. Formulas/cache are facts, never evaluated."""
    MAX_FILE=8*1024*1024
    MAX_TOTAL=64*1024*1024
    MAX_MEMBER=8*1024*1024
    MAX_RATIO=100
    MAX_ROWS=10000
    MAX_COLS=200
    MAX_CELLS=50000
    MAX_WORKBOOK_CELLS=100000
    MAX_RANGE_WORK=100000  # Per load; the two fixed loads have twice this bound.
    MAX_RANGES=1024
    MAX_TEXT=4096
    MAX_TEXT_TOTAL=1024*1024
    MAX_STRINGS=50000
    MAX_STYLES=10000  # XML style nodes, including nested font/fill/alignment nodes.

    def __init__(self, *, first_row_as_header: bool = False):
        if type(first_row_as_header) is not bool:
            raise ValueError('first_row_as_header must be a boolean')
        self.first_row_as_header = first_row_as_header

    def _bounds(self,area):
        try:a,b,c,d=range_boundaries(area)
        except (ValueError,TypeError) as exc:raise ParserError('XLSX_SHEET_LIMIT') from exc
        if not all(isinstance(n,int) and n>0 for n in (a,b,c,d)) or a>c or b>d or c>self.MAX_COLS or d>self.MAX_ROWS or c*d>self.MAX_CELLS:
            raise ParserError('XLSX_SHEET_LIMIT')
        return a,b,c,d

    def _preflight(self,path):
        if path.stat().st_size>self.MAX_FILE:raise ParserError('XLSX_ARCHIVE_LIMIT')
        sheets={};workbook_xml=None;relations_xml=None
        total_grid=0;range_work=0;text_total=0;shared_lengths=[];shared_refs={}

        def text_budget(length):
            nonlocal text_total
            if length>self.MAX_TEXT:raise ParserError('XLSX_TEXT_LIMIT')
            text_total+=length
            if text_total>self.MAX_TEXT_TOTAL:raise ParserError('XLSX_TEXT_LIMIT')

        def string_length(node):
            # openpyxl Text.content joins plain text and rich runs, excluding rPh.
            return sum(len(t.text or '') for t in node.findall(NS+'t'))+sum(
                len(t.text or '') for run in node.findall(NS+'r') for t in run.findall(NS+'t'))
        with zipfile.ZipFile(path) as archive:
            infos=archive.infolist()
            if len(infos)>2000 or sum(i.file_size for i in infos)>self.MAX_TOTAL:raise ParserError('XLSX_ARCHIVE_LIMIT')
            if len({i.filename for i in infos})!=len(infos):raise ParserError('XLSX_ARCHIVE_INVALID')
            for info in infos:
                name=info.filename
                if info.flag_bits&1 or info.file_size>self.MAX_MEMBER or info.file_size/max(1,info.compress_size)>self.MAX_RATIO:
                    raise ParserError('XLSX_ARCHIVE_LIMIT')
                if 'vbaproject' in name.lower() or name.startswith('xl/externalLinks/'):
                    raise ParserError('XLSX_ACTIVE_CONTENT_UNSUPPORTED')
                if not (name.endswith('.xml') or name.endswith('.rels')):continue
                data=archive.read(name)
                try:tree=fromstring(data,forbid_dtd=True,forbid_entities=True,forbid_external=True)
                except (DefusedXmlException,ParseError) as exc:raise ParserError('XLSX_UNSAFE_XML') from exc
                if name=='xl/workbook.xml':workbook_xml=tree
                if name=='xl/_rels/workbook.xml.rels':relations_xml=tree
                if tree.tag==NS+'sst':
                    strings=tree.findall(NS+'si')
                    if len(strings)>self.MAX_STRINGS:raise ParserError('XLSX_TEXT_LIMIT')
                    for item in strings:
                        length=string_length(item);text_budget(length);shared_lengths.append(length)
                    # Phonetic text is also bounded though not part of Text.content.
                    for node in tree.iter(NS+'rPh'):
                        for text in node.iter(NS+'t'):text_budget(len(text.text or ''))
                if tree.tag==NS+'styleSheet':
                    if sum(1 for _ in tree.iter())>self.MAX_STYLES:raise ParserError('XLSX_STYLE_LIMIT')
                    for node in tree.iter():
                        if len(node.attrib.get('formatCode',''))>self.MAX_TEXT:raise ParserError('XLSX_STYLE_LIMIT')
                if tree.tag==NS+'worksheet':
                    if len(sheets)>=32:raise ParserError('XLSX_SHEET_LIMIT')
                    dimension=tree.find(NS+'dimension')
                    if dimension is not None:self._bounds(dimension.attrib.get('ref',''))
                    raw={};max_row=1;max_col=1;cell_count=0;row_index=0
                    data_node=tree.find(NS+'sheetData')
                    for row in (() if data_node is None else data_node.findall(NS+'row')):
                        row_index=int(row.attrib.get('r',row_index+1))
                        self._bounds(f'A{row_index}')
                        column=0
                        for cell in row.findall(NS+'c'):
                            cell_count+=1
                            if cell_count>self.MAX_CELLS:raise ParserError('XLSX_SHEET_LIMIT')
                            coord=cell.attrib.get('r')
                            if coord:
                                if not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]*',coord):raise ParserError('XLSX_SHEET_LIMIT')
                                column,cell_row,_,_=self._bounds(coord)
                            else:
                                column+=1;cell_row=row_index
                                coord=f'{get_column_letter(column)}{cell_row}';self._bounds(coord)
                            max_row=max(max_row,cell_row);max_col=max(max_col,column)
                            formula=cell.find(NS+'f');value=cell.find(NS+'v');inline=cell.find(NS+'is')
                            if formula is not None:
                                text_budget(len(formula.text or ''))
                                # Array/shared formula endpoints must not hide sparse extents.
                                if formula.attrib.get('ref'):self._bounds(formula.attrib['ref'])
                            if inline is not None:text_budget(string_length(inline))
                            if cell.attrib.get('t')=='s' and value is not None:
                                index=int(value.text or '-1')
                                if index<0:raise ParserError('XLSX_ARCHIVE_INVALID')
                                shared_refs[index]=shared_refs.get(index,0)+1
                            elif value is not None:text_budget(len(value.text or ''))
                            if cell.attrib.get('t','n')=='n' and formula is None and value is not None:raw[coord]=value.text
                    merges=[];ranges=0
                    for tag in ('mergeCell','hyperlink'):
                        for node in tree.iter(NS+tag):
                            ranges+=1
                            if ranges>self.MAX_RANGES:raise ParserError('XLSX_SHEET_LIMIT')
                            a,b,c,d=self._bounds(node.attrib.get('ref',''))
                            if tag=='mergeCell':
                                # Bounded endpoint comparisons; never materialize a range.
                                if any(a<=x2 and x1<=c and b<=y2 and y1<=d for x1,y1,x2,y2 in merges):
                                    raise ParserError('XLSX_SHEET_LIMIT')
                                merges.append((a,b,c,d))
                            range_work+=(c-a+1)*(d-b+1)
                            if range_work>self.MAX_RANGE_WORK:raise ParserError('XLSX_SHEET_LIMIT')
                            max_row=max(max_row,d);max_col=max(max_col,c)
                    self._bounds(f'A1:{get_column_letter(max_col)}{max_row}')
                    total_grid+=max_row*max_col
                    if total_grid>self.MAX_WORKBOOK_CELLS:raise ParserError('XLSX_SHEET_LIMIT')
                    sheets[name]=raw
        for index,count in shared_refs.items():
            if index>=len(shared_lengths):raise ParserError('XLSX_ARCHIVE_INVALID')
            text_total+=shared_lengths[index]*count
            if text_total>self.MAX_TEXT_TOTAL:raise ParserError('XLSX_TEXT_LIMIT')
        if workbook_xml is None or relations_xml is None:raise ParserError('XLSX_ARCHIVE_INVALID')
        targets={node.attrib['Id']:node.attrib for node in relations_xml}
        named={};seen_targets=set();sheet_count=0
        for sheet in workbook_xml.iter(NS+'sheet'):
            sheet_count+=1
            if sheet_count>32:raise ParserError('XLSX_SHEET_LIMIT')
            relation=targets[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            if relation.get('TargetMode')=='External':raise ParserError('XLSX_ACTIVE_CONTENT_UNSUPPORTED')
            target=relation['Target']
            path_name=posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/'+target)
            if path_name not in sheets:raise ParserError('XLSX_ARCHIVE_INVALID')
            if path_name in seen_targets or sheet.attrib['name'] in named:raise ParserError('XLSX_ARCHIVE_INVALID')
            seen_targets.add(path_name)
            named[sheet.attrib['name']]=sheets[path_name]
        if len(named)>32:raise ParserError('XLSX_SHEET_LIMIT')
        return named

    @staticmethod
    def _value(value):
        if isinstance(value,(datetime,date,time)):return value.isoformat(),type(value).__name__
        if value is None:return None,'empty'
        if isinstance(value,bool):return value,'boolean'
        if isinstance(value,(int,float)):return value,'number'
        return str(value),'text'

    @staticmethod
    def _display(value,kind,number_format):
        if value is None:return 'EMPTY'
        if kind in {'datetime','date','time'}:return str(value)
        if kind=='boolean':return 'TRUE' if value else 'FALSE'
        if kind=='number' and re.fullmatch(r'0(?:\.(0+))?%',number_format):
            match=re.fullmatch(r'0(?:\.(0+))?%',number_format)
            return f'{value*100:.{len(match.group(1) or "")}f}%'
        return str(value)

    def parse(self,path:Path,document_id:str,version_id:str)->NormalizedDocument:
        books=[]
        try:
            if not xml_functions.DEFUSEDXML:raise ParserError('XLSX_XML_PROTECTION_UNAVAILABLE')
            raw_sheets=self._preflight(path)
            # Storage keys have no extension. Eager loads accept binary streams;
            # each context closes its handle, including a failed second load.
            for cached in (False,True):
                with path.open('rb') as source:
                    books.append(openpyxl.load_workbook(source,data_only=cached,read_only=False,keep_vba=False,keep_links=False))
            formula_book,cache_book=books
            if len(formula_book.worksheets)>32:raise ParserError('XLSX_SHEET_LIMIT')
            tables=[];blocks=[];sections=[];locators=[];pieces=[];offset=0;total_cells=0
            for index,ws in enumerate(formula_book.worksheets):
                self._bounds(ws.calculate_dimension());total_cells+=ws.max_row*ws.max_column
                if total_cells>self.MAX_WORKBOOK_CELLS:raise ParserError('XLSX_SHEET_LIMIT')
                cached_ws=cache_book[ws.title];raw=raw_sheets[ws.title]
                merged=tuple(str(r) for r in ws.merged_cells.ranges)
                anchors={}
                for area in merged:
                    a,b,c,d=range_boundaries(area)
                    for row in range(b,d+1):
                        for col in range(a,c+1):anchors[(row,col)]=f'{get_column_letter(a)}{b}'
                headers=(1,) if self.first_row_as_header and ws.max_row >= 2 else ()
                if len(headers)>8:raise ParserError('XLSX_TEXT_LIMIT')
                labels_by_column={};counts={}
                for col in range(1,ws.max_column+1):
                    anchor=anchors.get((1,col));head=ws[anchor] if anchor else ws.cell(1,col)
                    label=str(head.value).strip() if headers and head.value is not None and head.data_type!='f' else get_column_letter(col)
                    label=label or get_column_letter(col)
                    if len(label)>512:raise ParserError('XLSX_TEXT_LIMIT')
                    counts[label]=counts.get(label,0)+1
                    labels_by_column[col]=label if counts[label]==1 else f'{label}_{counts[label]}'
                cells=[]
                for row in range(1,ws.max_row+1):
                    for col in range(1,ws.max_column+1):
                        cell=ws.cell(row,col);coordinate=cell.coordinate;is_formula=cell.data_type=='f'
                        value,kind=self._value(None if is_formula else cell.value)
                        cached,cache_kind=self._value(cached_ws.cell(row,col).value if is_formula else None)
                        formula=str(cell.value) if is_formula else None
                        display=(f'FORMULA={formula}; CACHED='+('UNKNOWN_CACHE_MISSING' if cached is None else self._display(cached,cache_kind,cell.number_format))) if formula else self._display(value,kind,cell.number_format)
                        labels=[labels_by_column[col]]
                        lexical=raw.get(coordinate)
                        # Pair-load values are authoritative. A lexical candidate
                        # must agree before it is attributed to this named sheet.
                        if lexical is not None:
                            try:
                                if kind!='number' or float(lexical)!=float(value):lexical=None
                            except ValueError:lexical=None
                        cells.append(TableCell(coordinate=coordinate,row=row,column=col,value=value,value_type=kind,display=display,number_format=cell.number_format,raw_number=lexical,formula=formula,cached_value=cached,cache_status=('missing' if cached is None else 'present') if formula else 'not_applicable',merged_anchor=anchors.get((row,col)),column_headers=tuple(labels)))
                table_id=f'sheet-{index+1}'
                table=DocumentTable(table_id=table_id,sheet=ws.title,cell_range=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}',header_rows=headers,header_detection='explicit_first_row' if headers else 'column_letters',row_representation='key_value',merged_ranges=merged,cells=cells,start=offset,end=offset)
                rows=table_row_texts(table);text=''.join(content for _,content in rows)
                table=table.model_copy(update={'end':offset+len(text)})
                if not text:continue
                tables.append(table);pieces.append(text)
                blocks.append(DocumentBlock(block_id=table_id,kind='table',table_id=table_id,start=offset,end=offset+len(text)))
                sections.append(DocumentSection(section_id=table_id,heading=ws.title,heading_path=(ws.title,),level=1,start=offset,end=offset+len(text),content_type='table'))
                locators.append(SourceLocator(kind='table',sheet=ws.title,table_id=table_id,cell_range=table.cell_range,start=offset,end=offset+len(text)))
                offset+=len(text)
            text=''.join(pieces)
            if not any(c.value is not None or c.formula for table in tables for c in table.cells):raise ParserError('XLSX_EMPTY_TEXT')
            return NormalizedDocument(document_id=document_id,version_id=version_id,title=path.stem,media_type=MEDIA,markdown_content=text,sections=sections,source_locators=locators,tables=tables,blocks=blocks,content_sha256=hashlib.sha256(text.encode()).hexdigest(),parser_version='xlsx/openpyxl-3.1.5/row-v2',parser_engine='openpyxl',source_mapping_available=True)
        except ParserError:raise
        except (zipfile.BadZipFile,OSError,ValueError,KeyError,ParseError,DefusedXmlException) as exc:raise ParserError('XLSX_ARCHIVE_INVALID') from exc
        finally:
            for book in books:book.close()
