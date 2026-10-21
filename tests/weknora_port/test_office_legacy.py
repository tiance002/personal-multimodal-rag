"""Fixed Office behavior; tool processes are explicit substitutes, decoders are real."""
import base64
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from backend.weknora_port.ingestion import DocumentRequest, ParserRegistry, parse_document


def office_payloads():
    from docx import Document
    png = io.BytesIO()
    Image.new("RGB", (3, 2), "blue").save(png, format="PNG")
    image = png.getvalue()
    doc = Document()
    doc.add_paragraph("Blue widgets use protocol W-17.")
    doc.add_picture(io.BytesIO(image))
    out = io.BytesIO(); doc.save(out)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(image), 0, 0)
    ppt = io.BytesIO(); deck.save(ppt)
    return out.getvalue(), ppt.getvalue(), image


def test_default_doc_reader_is_registered():
    registry = ParserRegistry()
    assert registry.resolve("", "legacy.doc") == "doc"
    assert registry.resolve("builtin", "legacy.doc") == "builtin"


def test_pptx_unresolved_image_survives_markitdown_retry(monkeypatch):
    from markitdown import MarkItDown
    _, pptx, png = office_payloads()
    calls = []
    def convert(self, stream, **kwargs):
        calls.append(kwargs["keep_data_uris"])
        if kwargs["keep_data_uris"]:
            raise ValueError("fixture converter cannot inline media")
        return SimpleNamespace(text_content="Blue widgets W-17. ![diagram](image1.png)")
    monkeypatch.setattr(MarkItDown, "convert", convert)
    engine, parsed = parse_document(DocumentRequest("slides.pptx", pptx, source_version="original-v1"))
    assert parsed.ok and engine == "markitdown"
    assert len(parsed.assets) == 1
    asset = parsed.assets[0]
    assert asset.path in parsed.content and asset.data == png
    assert parsed.metadata["source_version"] == "original-v1"
    assert calls == [True, False]


@pytest.fixture
def doc_classes():
    from backend.weknora_port.ingestion import upstream_doc_parser as port
    from backend.weknora_port.ingestion import _office_runtime_adapter as bridge
    from backend.weknora_port.ingestion import upstream_office_tempfile as temp
    from backend.weknora_port.ingestion._docx_runtime_adapter import DocumentModel
    from tests.weknora_port.test_ingest import _upstream_module, _upstream_package_stubs, UPSTREAM
    import types
    modules = _upstream_package_stubs()
    config = types.ModuleType("docreader.config"); config.CONFIG = bridge.CONFIG
    docx2 = types.ModuleType("docreader.parser.docx2_parser"); docx2.Docx2Parser = bridge.Docx2Parser
    document = types.ModuleType("docreader.models.document"); document.Document = DocumentModel
    modules.update({"docreader.config": config, "docreader.parser.docx2_parser": docx2,
                    "docreader.models.document": document, "docreader.utils.tempfile": temp,
                    "textract": types.ModuleType("textract")})
    with _upstream_module("docreader.parser.doc_parser", UPSTREAM / "docreader/parser/doc_parser.py", modules) as source:
        yield port, source


def install_doc_tools(monkeypatch, cls, converted, *, mode="success"):
    trace = []
    monkeypatch.setattr(cls, "_try_find_soffice", lambda self: None if mode == "missing" else "fixture-soffice")
    monkeypatch.setattr(cls, "_try_find_antiword", lambda self: "fixture-antiword")
    def execute(self, command):
        if command[0] == "fixture-antiword":
            trace.append(("antiword", Path(command[1]).read_bytes()))
            return b"Blue widgets use protocol W-17.", b"", 0
        profile = next(part for part in command if part.startswith("-env:UserInstallation="))
        out = Path(command[command.index("--outdir") + 1])
        trace.append(("soffice", profile, str(out), Path(command[-1]).read_bytes()))
        attempts = sum(row[0] == "soffice" for row in trace)
        if mode == "timeout": raise RuntimeError("fixture conversion timeout")
        if mode == "fail": return b"", b"fixture failure", 1
        if mode == "retry" and attempts < 3: return b"", b"fixture retry", 1
        if mode != "no-output": (out / "converted.docx").write_bytes(converted)
        return b"", b"", 0
    # Both implementations execute their own conversion/retry bodies; only the
    # external command is replaced. Real converted DOCX bytes are decoded below.
    monkeypatch.setattr(cls.__module__ and __import__(cls.__module__, fromlist=["SandboxExecutor"]).SandboxExecutor,
                        "execute_in_sandbox", execute)
    return trace


@pytest.mark.parametrize("mode", ["success", "retry", "fail", "no-output", "missing", "timeout"])
def test_doc_conversion_and_antiword_same_fixed_source(monkeypatch, doc_classes, mode):
    import re
    port, source = doc_classes
    docx, _, _ = office_payloads()
    payload = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"synthetic legacy tool input"
    results = []
    for module in (port, source):
        with monkeypatch.context() as patch:
            patch.setattr(module.time, "sleep", lambda _seconds: None)
            trace = install_doc_tools(patch, module.DocParser, docx, mode=mode)
            parsed = module.DocParser(file_name="legacy.doc").parse_into_text(payload)
            profiles = [row[1] for row in trace if row[0] == "soffice"]
            assert len(profiles) == len(set(profiles))
            assert all(not Path(row[2]).exists() for row in trace if row[0] == "soffice")
            normalized = re.sub(r"images/[^)]+", "images/IMAGE", parsed.content)
            results.append((normalized, sorted(base64.b64decode(v) for v in parsed.images.values()), [row[0] for row in trace]))
    assert results[0] == results[1]
    if mode in ("fail", "no-output"): assert results[0][2] == ["soffice"] * 3 + ["antiword"]
    elif mode == "retry": assert results[0][2] == ["soffice"] * 3
    elif mode == "timeout": assert results[0][2] == ["soffice", "antiword"]
    elif mode == "missing": assert results[0][2] == ["antiword"]
    if mode in ("success", "retry"): assert results[0][1]


def test_doc_no_tools_and_empty_first_parser_match_source(monkeypatch, doc_classes):
    from backend.weknora_port.ingestion._docx_runtime_adapter import DocumentModel
    for module in doc_classes:
        with monkeypatch.context() as patch:
            patch.setattr(module.DocParser, "_try_find_soffice", lambda self: None)
            patch.setattr(module.DocParser, "_try_find_antiword", lambda self: None)
            assert module.DocParser(file_name="no-tools.doc").parse_into_text(b"fixture").content == ""
        parser = module.DocParser(file_name="empty.doc")
        parser._parsers = [SimpleNamespace(parse_into_text=lambda _raw: DocumentModel())]
        with monkeypatch.context() as patch:
            patch.setattr(parser, "_try_convert_doc_to_docx", lambda _path: b"converted")
            patch.setattr(parser, "_parse_with_antiword", lambda _path: pytest.fail("source does not fallback on a truthy empty Document"))
            assert parser.parse_into_text(b"fixture").content == ""


@pytest.mark.parametrize("mode", ["retry", "fail", "no-output", "missing", "timeout", "oserror"])
def test_ppt_converter_same_source_retry_profiles_and_failure(monkeypatch, mode):
    import importlib.util
    from backend.weknora_port.ingestion import upstream_ppt_convert as port
    from tests.weknora_port.test_ingest import _upstream_module, _upstream_package_stubs, UPSTREAM
    import types
    excel = types.ModuleType("docreader.parser.excel_convert"); excel.find_soffice = lambda: None
    modules = _upstream_package_stubs(); modules["docreader.parser.excel_convert"] = excel
    _, pptx, _ = office_payloads()
    outcomes = []
    with _upstream_module("docreader.parser.ppt_convert", UPSTREAM / "docreader/parser/ppt_convert.py", modules) as source:
        for module in (port, source):
            trace = []
            with monkeypatch.context() as patch:
                patch.setattr(module, "find_soffice", lambda: None if mode == "missing" else "fixture-soffice")
                patch.setattr(module.time, "sleep", lambda _seconds: None)
                def run(command, **kwargs):
                    assert kwargs["timeout"] == 120
                    trace.append(command)
                    if mode == "timeout": raise module.subprocess.TimeoutExpired(command, 120)
                    if mode == "oserror": raise OSError("fixture start failure")
                    out = Path(command[command.index("--outdir") + 1])
                    if mode == "retry" and len(trace) == 3: (out / "out.pptx").write_bytes(pptx)
                    return SimpleNamespace(returncode=1 if mode == "fail" else 0, stderr=b"fixture")
                patch.setattr(module.subprocess, "run", run)
                result = module.convert_ppt_to_pptx_bytes(b"legacy")
                profiles = [next(v for v in cmd if v.startswith("-env:UserInstallation=")) for cmd in trace]
                assert len(profiles) == len(set(profiles))
                outcomes.append((result, len(trace)))
    assert outcomes[0] == outcomes[1]
    assert outcomes[0][1] == (0 if mode == "missing" else 1 if mode in ("timeout", "oserror") else 3)


def test_ppt_normalize_real_zip_and_failure_isolation(monkeypatch):
    from backend.weknora_port.ingestion import upstream_ppt_convert as port
    _, pptx, _ = office_payloads()
    assert port.normalize_ppt_bytes(pptx, "ppt") == (pptx, ".pptx")
    monkeypatch.setattr(port, "find_soffice", lambda: None)
    engine, parsed = parse_document(DocumentRequest("legacy.ppt", b"legacy binary bytes"))
    assert engine == "markitdown" and not parsed.ok
    assert parsed.failures[0].code == "PARSER_FAILED"
    assert "LibreOffice" in parsed.failures[0].message
    _, healthy = parse_document(DocumentRequest("still-works.txt", b"healthy fixture"))
    assert healthy.ok and healthy.content == "healthy fixture"


def test_media_source_oracle_keeps_attached_and_data_uri_images(monkeypatch):
    import re
    from markitdown import MarkItDown
    from tests.weknora_port.test_ingest import _upstream_markitdown
    _, pptx, png = office_payloads()
    encoded = base64.b64encode(png).decode()
    text = "![embedded](local.png) ![inline](data:image/png;base64," + encoded + ")"
    monkeypatch.setattr(MarkItDown, "convert", lambda *_args, **_kwargs: SimpleNamespace(text_content=text))
    expected, expected_images = _upstream_markitdown("slides.pptx", pptx)
    _, parsed = parse_document(DocumentRequest("slides.pptx", pptx))
    normalize = lambda value: re.sub(r"images/[^)]+", "images/IMAGE", value)
    assert normalize(parsed.content) == normalize(expected)
    assert sorted(a.data for a in parsed.assets) == sorted(expected_images)
    assert len(parsed.assets) == 2 and all(a.path in parsed.content for a in parsed.assets)


@pytest.mark.parametrize("mode", ["success", "timeout"])
def test_doc_runner_real_child_process_and_bounded_cleanup(monkeypatch, doc_classes, mode):
    import subprocess
    import sys
    for module in doc_classes:
        processes = []
        original = subprocess.Popen
        def capture(*args, **kwargs):
            process = original(*args, **kwargs)
            processes.append(process)
            return process
        with monkeypatch.context() as patch:
            patch.setattr(module.subprocess, "Popen", capture)
            runner = module.SandboxExecutor(proxy="http://synthetic.invalid:1", default_timeout=.1 if mode == "timeout" else 5)
            code = "import time; time.sleep(30)" if mode == "timeout" else "import os; print(os.environ['HTTPS_PROXY'])"
            if mode == "timeout":
                with pytest.raises(RuntimeError, match="timeout"):
                    runner._execute_with_proxy([sys.executable, "-B", "-c", code])
                assert processes[0].poll() is not None
                assert processes[0].stdout.closed and processes[0].stderr.closed
            else:
                output, error, status = runner._execute_with_proxy([sys.executable, "-B", "-c", code])
                assert status == 0 and output.strip() == b"http://synthetic.invalid:1" and not error


@pytest.mark.parametrize("primary", ["", "   "])
def test_doc_proxy_present_primary_overrides_fallback_verbatim(monkeypatch, primary):
    from backend.weknora_port.ingestion._office_runtime_adapter import _proxy_setting
    monkeypatch.setenv("DOCREADER_EXTERNAL_HTTPS_PROXY", primary)
    monkeypatch.setenv("EXTERNAL_HTTPS_PROXY", "http://fallback.invalid:1")
    import ast
    from tests.weknora_port.test_ingest import UPSTREAM
    config = ast.parse((UPSTREAM / "docreader/config.py").read_text(encoding="utf-8"))
    getters = [node for node in config.body if isinstance(node, ast.FunctionDef) and node.name in ("_get_first_env", "_get_str")]
    import os
    from typing import Iterable, Optional, Tuple
    namespace = {"os": os, "Iterable": Iterable, "Optional": Optional, "Tuple": Tuple}
    exec(compile(ast.Module(body=getters, type_ignores=[]), "<fixed-source-config-getter>", "exec"), namespace)
    assert _proxy_setting() == namespace["_get_str"](["DOCREADER_EXTERNAL_HTTPS_PROXY", "EXTERNAL_HTTPS_PROXY"], "") == primary


def test_doc_bridge_preserves_explicit_file_type_and_filename(monkeypatch):
    from backend.weknora_port.ingestion._office_runtime_adapter import MarkitdownParser
    from backend.weknora_port.ingestion import upstream_docx_merge
    from markitdown import MarkItDown
    docx, _, _ = office_payloads()
    hints = []; merges = []
    monkeypatch.setattr(MarkItDown, "convert", lambda self, stream, **kw: (hints.append(kw["file_extension"]) or SimpleNamespace(text_content="fixture")))
    original = upstream_docx_merge.fill_vertical_merged_cells_docx
    monkeypatch.setattr(upstream_docx_merge, "fill_vertical_merged_cells_docx", lambda data: (merges.append(data) or original(data)))
    assert MarkitdownParser(file_name="input.bin", file_type="doc").parse_into_text(docx).content == "fixture"
    assert MarkitdownParser(file_name="legacy.doc", file_type="docx").parse_into_text(docx).content == "fixture"
    assert hints == [".doc", ".docx"]
    assert merges == [docx]
