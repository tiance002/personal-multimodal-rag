"""Actual E-reader Office libraries and scoped original upload/index/readback."""
import io,json,hashlib
from pathlib import Path
from docx import Document
from pptx import Presentation
from markitdown import MarkItDown
from importlib.metadata import version
from backend.weknora_port.ingestion import DocumentRequest,parse_document
from tests.weknora_port.native_fixture import isolated_mock_application
from tests.weknora_port.test_memory_storage_pg import actual_pg_run
from tests.weknora_port.test_ingest import _upstream_markitdown


def test_actual_docx_pptx_csv_source_pair_and_native_entry():
    doc=Document();doc.add_heading('Source office fixture',0);doc.add_paragraph('Blue widgets use protocol W-17.');output=io.BytesIO();doc.save(output)
    docx=output.getvalue()
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Source slides';slide.placeholders[1].text='Blue widgets use protocol W-17.'
    output=io.BytesIO();deck.save(output);pptx=output.getvalue()
    cases=[('office.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document',docx),('slides.pptx','application/vnd.openxmlformats-officedocument.presentationml.presentation',pptx),('table.csv','text/csv',b'Name,Protocol\nBlue widgets,W-17\n')]
    prepared=[]
    for name,media,raw in cases:
        expected,images=_upstream_markitdown(name,raw)
        engine,parsed=parse_document(DocumentRequest(name,raw))
        assert parsed.content==expected and not images
        prepared.append((name,media,raw,parsed.content,engine))
    async def scenario():
        async with isolated_mock_application() as (app,wire,root):
            kb=(await app.create_knowledge_base(app.principal,'Native Office fixture'))['id']
            await app.update_knowledge_base_consent(app.principal,kb,True)
            receipts=[]
            for name,media,raw,text,engine in prepared:
                result=await app.upload_document(app.principal,kb,name,media,raw);assert result['status']=='completed',result
                actual=await app.read_document(app.principal,kb,result['id']);assert actual.content==raw
                async with app.leases.lease() as guard:
                    rows=await (await guard.connection.execute("SELECT content,source_content FROM chunks WHERE tenant_id=%s AND knowledge_id=%s AND chunk_type='text'",(app.principal.tenant_id,result['id']))).fetchall()
                    assert rows and any((row['source_content'] or row['content'])==text for row in rows)
                    indexes=(await (await guard.connection.execute('SELECT count(*) AS n FROM embeddings WHERE knowledge_id=%s AND is_enabled=true',(result['id'],))).fetchone())['n']
                    assert indexes>0
                receipts.append({'format':name,'parser_engine':engine,'document':result['id'],'raw_sha256':hashlib.sha256(raw).hexdigest(),'source_text_sha256':hashlib.sha256(text.encode()).hexdigest(),'active_indices':indexes,'original_readback_equal':True,'source_text_equal':True})
            import mammoth,markitdown,pptx as slides
            assert str(Path('E:/WK/deps/reader')).lower() in str(Path(mammoth.__file__)).lower()
            Path('docs/weknora-port/evidence/P2_OFFICE_NATIVE_ENTRY_20261014.json').write_text(json.dumps({'kind':'REAL_OFFICE_PYTHON_LIBRARIES_FIXED_SOURCE_FIXTURE_PLUS_NATIVE_PG_ENTRY','reader_root':'E:/WK/deps/reader',
                'versions':{name:version(name) for name in ('mammoth','markitdown','python-pptx','pandas','numpy')},'cases':receipts,'storage_root':str(root),'mock_calls':len(wire),'real_provider_calls':0,'docreader_service_started':False,'old_environment_modified':False,
                'scope':'Simple DOCX/PPTX/CSV source text fixtures; complex PPTX media attachment/legacy DOC/PPT conversion/full format parity not claimed'},indent=2),encoding='utf-8')
    actual_pg_run(scenario())
