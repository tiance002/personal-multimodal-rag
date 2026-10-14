"""Actual fixed Python source Excel oracle and native upload/PG entry."""
import io,sys,types,importlib,contextlib,zipfile,json,hashlib
from pathlib import Path
from openpyxl import Workbook
from backend.weknora_port.ingestion import DocumentRequest,parse_document
from backend.weknora_port.ingestion._excel_runtime_adapter import Chunk,Document
from backend.weknora_port.ingestion._docx_runtime_adapter import BaseParser
from tests.weknora_port.native_fixture import isolated_mock_application
from tests.weknora_port.test_memory_storage_pg import actual_pg_run
from scripts.weknora_port_six_mock import prepare_public_fixture


@contextlib.contextmanager
def upstream_excel():
    root=Path('E:/RAG-reference/WeKnora-3e8b0bfc/docreader')
    saved={key:value for key,value in sys.modules.items() if key=='docreader' or key.startswith('docreader.')}
    for key in saved:sys.modules.pop(key)
    package=types.ModuleType('docreader');package.__path__=[str(root)]
    parser=types.ModuleType('docreader.parser');parser.__path__=[str(root/'parser')]
    models=types.ModuleType('docreader.models');models.__path__=[]
    dto=types.ModuleType('docreader.models.document');dto.Chunk=Chunk;dto.Document=Document
    base=types.ModuleType('docreader.parser.base_parser');base.BaseParser=BaseParser
    sys.modules.update({'docreader':package,'docreader.parser':parser,'docreader.models':models,'docreader.models.document':dto,'docreader.parser.base_parser':base})
    try:yield importlib.import_module('docreader.parser.excel_parser')
    finally:
        for key in list(sys.modules):
            if key=='docreader' or key.startswith('docreader.'):sys.modules.pop(key)
        sys.modules.update(saved)


def workbook_bytes():
    workbook=Workbook();sheet=workbook.active;sheet.title='Merged'
    sheet.append(['Group','Value','Value']);sheet.append(['Blue widgets',17,3.5]);sheet.append([None,18,None]);sheet.merge_cells('A2:A3')
    sheet.append([None,None,None]);sheet.append(['=_xlfn.IMAGE("example")','Actual text',None])
    second=workbook.create_sheet('Other');second.append(['Label','Count']);second.append(['W-17',2])
    output=io.BytesIO();workbook.save(output);return output.getvalue()


def dangling_sharedstrings(raw):
    source=zipfile.ZipFile(io.BytesIO(raw));target=io.BytesIO()
    with zipfile.ZipFile(target,'w') as result:
        for name in source.namelist():
            content=source.read(name)
            if name=='[Content_Types].xml':content=content.replace(b'</Types>',b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            result.writestr(name,content)
    return target.getvalue()


def test_fixed_excel_source_real_helpers_and_native_upload():
    raw=workbook_bytes();cases=[('default',raw,False),('header',raw,True),('repair',dangling_sharedstrings(raw),True)]
    receipts=[]
    for name,payload,header in cases:
        with upstream_excel() as source:
            expected=source.ExcelParser(file_name='source.xlsx',file_type='xlsx',xlsx_first_row_as_header=header).parse_into_text(payload)
        engine,parsed=parse_document(DocumentRequest('source.xlsx',payload,upload_parser_overrides={'xlsx_first_row_as_header':str(header)}))
        assert engine=='excel' and parsed.content==expected.content
        assert parsed.metadata['source_row_chunks']==[{'seq':r.seq,'start':r.start,'end':r.end} for r in expected.chunks]
        assert parsed.metadata['sheet_count']==2
        assert parsed.content.count('Blue widgets')==2
        assert 'IMAGE(' not in parsed.content
        receipts.append({'name':name,'raw_sha256':hashlib.sha256(payload).hexdigest(),'parsed_sha256':hashlib.sha256(parsed.content.encode()).hexdigest(),'rows':len(expected.chunks),'sheet_count':2})
    async def scenario():
        async with isolated_mock_application() as (app,wire,root):
            kb=(await app.create_knowledge_base(app.principal,'Source Excel fixture'))['id']
            await prepare_public_fixture(app,kb,'excel')
            document=await app.upload_document(app.principal,kb,'merged.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',raw)
            assert document['status']=='completed',document
            original=await app.read_document(app.principal,kb,document['id']);assert original.content==raw
            expected=parse_document(DocumentRequest('merged.xlsx',raw))[1].content
            async with app.leases.lease() as guard:
                rows=await (await guard.connection.execute("SELECT content,source_content FROM chunks WHERE tenant_id=%s AND knowledge_id=%s AND chunk_type='text' ORDER BY chunk_index",(app.principal.tenant_id,document['id']))).fetchall()
                assert rows and any((r['source_content'] or r['content'])==expected for r in rows)
                indexes=(await (await guard.connection.execute('SELECT count(*) AS n FROM embeddings WHERE knowledge_id=%s AND is_enabled=true',(document['id'],))).fetchone())['n']
                assert indexes>0
            Path('docs/weknora-port/evidence/P2_EXCEL_SOURCE_ENTRY_20261014.json').write_text(json.dumps({'kind':'FIXED_SOURCE_PYTHON_EXCEL_REAL_HELPERS_PLUS_NATIVE_UPLOAD_INDEX','cases':receipts,
                'document':document['id'],'storage_root':str(root),'active_indices':indexes,'original_hash_readback':True,'source_text_same':True,'mock_calls':len(wire),'real_provider_calls':0,'old_environment_changed':False},indent=2),encoding='utf-8')
    actual_pg_run(scenario())
