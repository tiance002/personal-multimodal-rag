from __future__ import annotations

import base64
import contextlib
import importlib.util
import io
import json
from pathlib import Path
import re
from concurrent.futures import ThreadPoolExecutor
import sys
import types

import fitz
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from backend.weknora_port.ingestion import ChunkConfig, DocumentRequest, ParserRegistry, ingest_document, parse_document, resolve_chunk_config
from backend.weknora_port.ingestion.core import _parse_docx_builtin, split, split_parent_child
from backend.weknora_port.ingestion.upstream_pdf_parser import PDFParser, _classify_page, _postprocess_pdf_text
from scripts.weknora_port_chunker_oracle import run_chunker_oracle


ROOT = Path(__file__).resolve().parents[2]
CHUNK_CASES = ROOT / "tests" / "weknora_port" / "fixtures" / "ingest_chunk_cases.json"
UPSTREAM = Path(r"E:\RAG-reference\WeKnora-3e8b0bfc")


@contextlib.contextmanager
def _upstream_module(module_name: str, source: Path, modules: dict[str, types.ModuleType]):
    """Load one fixed upstream module with only unrelated package edges stubbed."""
    previous = {name: sys.modules.get(name) for name in modules}
    sys.modules.update(modules)
    try:
        spec = importlib.util.spec_from_file_location(module_name, source)
        assert spec and spec.loader
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        yield module
    finally:
        sys.modules.pop(module_name, None)
        for name, old in previous.items():
            if old is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = old


def _upstream_package_stubs() -> dict[str, types.ModuleType]:
    package = types.ModuleType("docreader")
    package.__path__ = []
    models = types.ModuleType("docreader.models")
    models.__path__ = []
    parser = types.ModuleType("docreader.parser")
    parser.__path__ = []
    utils = types.ModuleType("docreader.utils")
    utils.__path__ = []
    config = types.ModuleType("docreader.config")
    config.CONFIG = types.SimpleNamespace(docx_max_pages=100000, markitdown_max_workers=1)
    document_module = types.ModuleType("docreader.models.document")

    class SourceDocument:
        def __init__(self, content="", images=None, chunks=None, metadata=None):
            self.content = content
            self.images = images or {}
            self.chunks = chunks or []
            self.metadata = metadata or {}

        def is_valid(self):
            return bool(self.content)

    class SourceChunk:
        def __init__(self, content="", seq=0, start=0, end=0, **_kwargs):
            self.content, self.seq, self.start, self.end = content, seq, start, end

    document_module.Document = SourceDocument
    document_module.Chunk = SourceChunk
    base_module = types.ModuleType("docreader.parser.base_parser")

    class SourceBaseParser:
        def __init__(self, file_name="", file_type=None, **_kwargs):
            self.file_name = file_name
            self.file_type = file_type or Path(file_name).suffix.lstrip(".")

    base_module.BaseParser = SourceBaseParser
    encode_module = types.ModuleType("docreader.utils.endecode")
    encode_module.encode_bytes = lambda text: text.encode("utf-8")
    encode_module.decode_image = lambda _image: ""
    utils.endecode = encode_module
    return {
        "docreader": package,
        "docreader.models": models,
        "docreader.models.document": document_module,
        "docreader.parser": parser,
        "docreader.parser.base_parser": base_module,
        "docreader.utils": utils,
        "docreader.utils.endecode": encode_module,
        "docreader.config": config,
    }


def _upstream_markitdown(filename: str, payload: bytes) -> tuple[str, list[bytes]]:
    stubbed = _upstream_package_stubs()
    chain = types.ModuleType("docreader.parser.chain_parser")
    chain.PipelineParser = type("PipelineParser", (stubbed["docreader.parser.base_parser"].BaseParser,), {})
    concurrency = types.ModuleType("docreader.parser.concurrency")
    concurrency.parser_worker_limit = lambda *_args: contextlib.nullcontext()
    docx_merge = types.ModuleType("docreader.parser.docx_merge")
    docx_merge.fill_vertical_merged_cells_docx = lambda data: data
    markdown_placeholder = types.ModuleType("docreader.parser.markdown_parser")
    markdown_placeholder.MarkdownParser = chain.PipelineParser
    ppt_convert = types.ModuleType("docreader.parser.ppt_convert")
    ppt_convert.normalize_ppt_bytes = lambda data, _ft: (data, ".pptx")
    pptx_media = types.ModuleType("docreader.parser.pptx_media")
    pptx_media.attach_pptx_media_to_markdown = lambda text, _data: (text, {})
    pptx_media.markdown_needs_pptx_media_attach = lambda _text: False
    stubbed.update({
        "docreader.parser.chain_parser": chain,
        "docreader.parser.concurrency": concurrency,
        "docreader.parser.docx_merge": docx_merge,
        "docreader.parser.markdown_parser": markdown_placeholder,
        "docreader.parser.ppt_convert": ppt_convert,
        "docreader.parser.pptx_media": pptx_media,
    })
    extension = Path(filename).suffix.lstrip(".").lower()
    source_path = UPSTREAM / "docreader" / "parser" / "markitdown_parser.py"
    with _upstream_module("docreader.parser.markitdown_parser", source_path, stubbed) as source:
        converted = source.StdMarkitdownParser(file_name=filename, file_type=extension).parse_into_text(payload)

    # Exercise the fixed upstream Markdown pipeline classes directly after the
    # real fixed-source MarkItDown wrapper.
    markdown_stubbed = _upstream_package_stubs()
    markdown_chain = types.ModuleType("docreader.parser.chain_parser")
    markdown_chain.PipelineParser = type("PipelineParser", (markdown_stubbed["docreader.parser.base_parser"].BaseParser,), {})
    markdown_stubbed["docreader.parser.chain_parser"] = markdown_chain
    endecode = markdown_stubbed["docreader.utils.endecode"]
    endecode.decode_bytes = lambda data: data.decode("utf-8", errors="replace")
    endecode.encode_image = lambda value, **_kwargs: base64.b64decode(value)
    with _upstream_module("docreader.parser.markdown_parser", UPSTREAM / "docreader" / "parser" / "markdown_parser.py", markdown_stubbed) as markdown:
        formatted = markdown.MarkdownTableFormatter().parse_into_text(converted.content.encode("utf-8"))
        extracted = markdown.MarkdownImageBase64().parse_into_text(formatted.content.encode("utf-8"))
    return extracted.content, [base64.b64decode(value) for value in extracted.images.values()]


def _png_bytes(color: tuple[int, int, int] = (12, 34, 56), size: tuple[int, int] = (256, 256)) -> bytes:
    image = Image.new("RGB", size, color)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def _pdf_bytes(*, scanned: bool, mixed: bool = False, text_overlay: bool = False) -> bytes:
    doc = fitz.open()
    text_page = doc.new_page(width=595, height=842)
    text_page.insert_text((72, 90), "Native text layer for extraction.")
    if scanned or mixed:
        image_page = doc.new_page(width=595, height=842)
        image_page.insert_image(image_page.rect, stream=_png_bytes())
        if text_overlay:
            image_page.insert_text((72, 90), "embedded OCR overlay")
    if scanned and not mixed:
        # Keep only the full-page image page for the scan-only fixture.
        doc.delete_page(0)
    result = doc.tobytes()
    doc.close()
    return result


def _repeated_header_pdf() -> bytes:
    doc = fitz.open()
    for page_number in range(4):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 55), "Document Running Header")
        page.insert_text((72, 180), f"Unique body content for page {page_number + 1}.")
        page.insert_text((72, 800), "2026")
    result = doc.tobytes()
    doc.close()
    return result


def _repeated_embedded_image_pdf() -> bytes:
    doc = fitz.open()
    logo = _png_bytes((90, 120, 150))
    for page_number in range(4):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 90), f"Native body page {page_number + 1} with enough text.")
        page.insert_image(fitz.Rect(100, 200, 228, 328), stream=logo)
    result = doc.tobytes()
    doc.close()
    return result


def _assert_adapter_matches_upstream_pdf(filename: str, data: bytes, overrides: dict[str, str] | None = None) -> None:
    upstream = PDFParser(file_name=filename, pdf_force_scanned=(overrides or {}).get("pdf_force_scanned")).parse_into_text(data)
    _, adapted = parse_document(DocumentRequest(filename, data, upload_parser_overrides=overrides or {}))
    assert adapted.content == upstream.content
    assert adapted.metadata["source_metadata"] == upstream.metadata
    assert [asset.data for asset in adapted.assets] == [base64.b64decode(value) for value in upstream.images.values()]


def _docx_bytes(image_data: bytes | None = None) -> bytes:
    doc = Document()
    doc.add_heading("Indexing Report", level=1)
    doc.add_paragraph("Opening paragraph before the table.")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    table.rows[0].cells[2].text = "Note"
    for index in range(120):
        cells = table.add_row().cells
        cells[0].text = f"row-{index}"
        cells[1].text = str(index)
        cells[2].text = "wide | cell"
    doc.add_paragraph("Picture follows.")
    doc.add_picture(io.BytesIO(image_data or _png_bytes()), width=Inches(0.5))
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "People"
    sheet.append(["Name", "Age", "Name", "Picture"])
    sheet.append(["Alice", 34, "A", '=_xlfn.DISPIMG("id-1",1)'])
    sheet.append([None, None, None])
    second = workbook.create_sheet("Next")
    second.append(["City", "Count"])
    second.append(["Oslo", 2])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(914400, 914400, 4572000, 914400)
    box.text_frame.text = "Synthetic deck title"
    slide.shapes.add_textbox(914400, 2200000, 4572000, 914400).text_frame.text = "One synthetic bullet"
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_registry_routes_native_and_scanned_pdf_pages_and_preserves_assets() -> None:
    name, native = parse_document(DocumentRequest("report.pdf", _pdf_bytes(scanned=False)))
    assert name == "pdf"
    _assert_adapter_matches_upstream_pdf("report.pdf", _pdf_bytes(scanned=False))
    assert native.metadata["page_count"] == 1
    assert native.metadata["text_page_count"] == 1
    assert native.metadata["scanned_page_count"] == 0
    assert "Native text layer" in native.content
    assert native.metadata["ocr_status"] == "NOT_RUN"

    _, scanned = parse_document(DocumentRequest("scan.pdf", _pdf_bytes(scanned=True)))
    _assert_adapter_matches_upstream_pdf("scan.pdf", _pdf_bytes(scanned=True))
    assert scanned.metadata["page_count"] == 1
    assert scanned.metadata["scanned_page_count"] == 1
    assert scanned.metadata["text_page_count"] == 0
    assert "scanned_image" in scanned.metadata["pages"][0]["route"]
    assert scanned.assets[0].source_type == "scanned_pdf"
    assert scanned.assets[0].data.startswith(b"\xff\xd8")
    assert scanned.metadata["ocr_status"] == "NOT_RUN"

    _, hybrid = parse_document(DocumentRequest("hybrid.pdf", _pdf_bytes(scanned=False, mixed=True)))
    _assert_adapter_matches_upstream_pdf("hybrid.pdf", _pdf_bytes(scanned=False, mixed=True))
    assert [page["route"] for page in hybrid.metadata["pages"]] == ["text_layer", "scanned_image"]

    # Image coverage remains primary even when a PDF carries a misleading OCR layer.
    _, overlay = parse_document(DocumentRequest("overlay.pdf", _pdf_bytes(scanned=True, text_overlay=True)))
    assert overlay.metadata["pages"][0]["route"] == "scanned_image"
    assert "embedded OCR overlay" not in overlay.content

    _, forced = parse_document(DocumentRequest(
        "forced.pdf", _pdf_bytes(scanned=False), upload_parser_overrides={"pdf_force_scanned": "true"}
    ))
    _assert_adapter_matches_upstream_pdf("forced.pdf", _pdf_bytes(scanned=False), {"pdf_force_scanned": "true"})
    assert forced.metadata["scanned_page_count"] == 1
    assert forced.metadata["ocr_status"] == "NOT_RUN"

    # A mismatched requested engine falls through to the extension default.
    engine, _ = parse_document(DocumentRequest("report.pdf", _pdf_bytes(scanned=False), parser_engine="excel"))
    assert engine == "pdf"


def test_pdfium_routes_are_serialized_across_concurrent_uploads() -> None:
    payloads = [_pdf_bytes(scanned=False), _pdf_bytes(scanned=True), _pdf_bytes(scanned=False, mixed=True)] * 2
    with ThreadPoolExecutor(max_workers=6) as pool:
        outputs = list(pool.map(lambda pair: PDFParser(file_name=f"parallel-{pair[0]}.pdf").parse_into_text(pair[1]), enumerate(payloads)))
    assert len(outputs) == len(payloads)
    assert all(doc.metadata["page_count"] >= 1 for doc in outputs)


def test_image_markdown_and_embedded_data_assets_preserve_content_and_failure_metadata() -> None:
    raw = _png_bytes()
    engine, image = parse_document(DocumentRequest("diagram.png", raw))
    assert engine == "image"
    assert image.content == "![diagram.png](images/diagram.png)"
    assert image.assets[0].data == raw
    assert image.metadata["image_source_type"] == "standalone_image"

    payload = base64.b64encode(raw).decode()
    _, markdown = parse_document(DocumentRequest(
        "guide.md", f"# Setup\n\n|Name|Age|\n|:---|---:|\n|Alice|3|\n\n![diagram](data:image/png;base64,{payload})".encode()
    ))
    assert markdown.content.startswith("# Setup")
    assert "data:image" not in markdown.content
    assert "| Name | Age |" in markdown.content
    assert "| :--- | ---: |" in markdown.content
    assert markdown.assets[0].data == raw
    assert markdown.assets[0].source_type == "markdown_data_uri"

    _, failed = parse_document(DocumentRequest("broken.pdf", b"not a pdf", document_id="d-1", source_version="v-2"))
    assert failed.failures and failed.failures[0].stage == "parse"
    assert failed.metadata["status"] == "failed"
    assert failed.metadata["filename"] == "broken.pdf"
    assert failed.metadata["input_bytes"] == len(b"not a pdf")


def test_pdf_source_cleanup_removes_repeated_headers_footers_and_watermark_images() -> None:
    _, repeated_text = parse_document(DocumentRequest("headers.pdf", _repeated_header_pdf()))
    assert repeated_text.metadata["page_count"] == 4
    assert "Document Running Header" not in repeated_text.content
    assert "2026" not in repeated_text.content
    assert "Unique body content for page 3." in repeated_text.content

    _, repeated_images = parse_document(DocumentRequest("logos.pdf", _repeated_embedded_image_pdf()))
    assert repeated_images.metadata["text_page_count"] == 4
    assert repeated_images.metadata["embedded_image_count"] == 0
    assert not [asset for asset in repeated_images.assets if asset.source_type == "embedded_pdf_image"]


def test_pdf_page_classification_thresholds_and_artifact_cleanup_match_upstream() -> None:
    assert _classify_page(0.5, 100) == "scanned"
    assert _classify_page(0.499, 10) == "text"
    assert _classify_page(0.1, 9) == "scanned"
    assert _classify_page(0.099, 0) == "text"

    page = fitz.open()
    pdf_page = page.new_page()
    pdf_page.insert_text((72, 90), "hy\ufffepen\n4\narXiv: 2401.12345")
    pdf_page.insert_text((72, 130), "re\ufffepair text")
    raw = page.tobytes()
    page.close()
    _, parsed = parse_document(DocumentRequest("artifacts.pdf", raw))
    assert "re\ufffepair" not in parsed.content
    assert "2401.12345" not in parsed.content
    assert "repair" in _postprocess_pdf_text("re\ufffepair text")


def test_docx_preserves_heading_table_headers_long_rows_and_embedded_image() -> None:
    payload = _docx_bytes()
    expected_content, expected_images = _upstream_markitdown("report.docx", payload)
    engine, parsed = parse_document(DocumentRequest("report.docx", payload))
    assert engine == "docx"
    stable_path = lambda text: re.sub(r"images/[0-9a-f-]{36}\.png", "images/<uuid>.png", text)
    assert stable_path(parsed.content) == stable_path(expected_content)
    assert parsed.content.startswith("# Indexing Report")
    assert "| Name | Value | Note |" in parsed.content
    assert "row-119" in parsed.content
    assert parsed.content.index("Opening paragraph") < parsed.content.index("| Name | Value")
    assert parsed.content.index("| row-119") < parsed.content.index("Picture follows")
    assert len(parsed.assets) == 1
    assert parsed.assets[0].source_type == "docx_markdown_data_uri"
    assert parsed.assets[0].data.startswith(b"\x89PNG")
    assert [asset.data for asset in parsed.assets] == expected_images
    assert parsed.metadata["docx_parser_route"] == "markitdown_primary"


def test_docx_failure_fallback_executes_fixed_upstream_process_pool_route(monkeypatch) -> None:
    from backend.weknora_port.ingestion import upstream_docx_parser as source

    payload = _docx_bytes()
    entered: list[tuple[int, int]] = []
    original = source.Docx._execute_multiprocess_tasks

    def observe_process_pool(self, args_list, max_workers):
        entered.append((len(args_list), max_workers))
        return original(self, args_list, max_workers)

    monkeypatch.setattr(source.Docx, "_execute_multiprocess_tasks", observe_process_pool)
    actual = _parse_docx_builtin(DocumentRequest("fallback.docx", payload))
    assert entered and entered[0][0] >= 1 and entered[0][1] >= 1
    expected = source.DocxParser(file_name="fallback.docx").parse_into_text(payload)
    assert actual.content == expected.content
    assert [asset.data for asset in actual.assets] == [base64.b64decode(value) for value in expected.images.values()]
    assert actual.metadata["docx_parser_route"] == "docx_source_fallback"


def test_excel_header_override_and_multisheet_rows() -> None:
    stubbed = _upstream_package_stubs()
    convert = types.ModuleType("docreader.parser.excel_convert")
    convert.detect_excel_format = lambda _data: "xlsx"
    convert.engine_for_format = lambda _ext: "openpyxl"
    convert.normalize_excel_bytes = lambda data, **_kwargs: data
    convert.convert_excel_to_xlsx_bytes = lambda data, **_kwargs: data
    repair = types.ModuleType("docreader.parser.xlsx_repair")
    repair.repair_xlsx_bytes = lambda _data: None
    merge = types.ModuleType("docreader.parser.xlsx_merge")
    merge.fill_merged_cells_xlsx = lambda data: data
    stubbed.update({
        "docreader.parser.excel_convert": convert,
        "docreader.parser.xlsx_repair": repair,
        "docreader.parser.xlsx_merge": merge,
    })
    payload = _xlsx_bytes()
    with _upstream_module("docreader.parser.excel_parser", UPSTREAM / "docreader" / "parser" / "excel_parser.py", stubbed) as source:
        expected = source.ExcelParser(file_name="people.xlsx", file_type="xlsx", xlsx_first_row_as_header=True).parse_into_text(payload)
    engine, parsed = parse_document(DocumentRequest(
        "people.xlsx", payload, upload_parser_overrides={"xlsx_first_row_as_header": "true"}
    ))
    assert engine == "excel"
    assert parsed.content == expected.content
    assert "Name: Alice,Age: 34,Name_2: A\n" in parsed.content
    assert "City: Oslo,Count: 2\n" in parsed.content
    assert "Name: Name" not in parsed.content
    assert parsed.metadata["sheet_count"] == 2
    assert parsed.metadata["row_count"] == 2
    assert parsed.metadata["row_sources"] == [{"sheet": "People", "row": 2}, {"sheet": "Next", "row": 2}]


def test_markitdown_registry_route_matches_fixed_wrapper_and_markdown_pipeline() -> None:
    payload = b"Name,Count\nAlice,3\nBob,5\n"
    expected_content, expected_images = _upstream_markitdown("people.csv", payload)
    engine, parsed = parse_document(DocumentRequest("people.csv", payload))
    assert engine == "markitdown"
    assert parsed.content == expected_content
    assert [asset.data for asset in parsed.assets] == expected_images

    pptx_payload = _pptx_bytes()
    pptx_expected, pptx_expected_images = _upstream_markitdown("deck.pptx", pptx_payload)
    pptx_engine, pptx_parsed = parse_document(DocumentRequest("deck.pptx", pptx_payload))
    assert pptx_engine == "markitdown"
    assert pptx_parsed.content == pptx_expected
    assert [asset.data for asset in pptx_parsed.assets] == pptx_expected_images


def test_config_priority_parent_identity_and_storage_neutral_records() -> None:
    cfg = resolve_chunk_config(
        {"chunk_size": 100, "chunk_overlap": 10, "enable_parent_child": True, "parent_chunk_size": 200, "child_chunk_size": 60},
        {"chunk_size": 150, "enable_parent_child": False},
        {"chunk_size": 0, "enable_parent_child": True, "strategy": "recursive"},
    )
    assert cfg.chunk_size == 150  # zero upload size is absent; upload bool is explicit
    assert cfg.enable_parent_child is True
    assert cfg.parent_chunk_size == 200 and cfg.child_chunk_size == 60

    exact = split_parent_child("A small source paragraph.", ChunkConfig(enable_parent_child=True, parent_chunk_size=256, child_chunk_size=256))
    assert len(exact.parents) == 0
    assert exact.children[0].parent_index == -1

    request = DocumentRequest(
        "notes.txt", ("A source paragraph. " * 40).encode(), owner_id="owner-1", knowledge_base_id="kb-1",
        document_id="doc-1", source_version="v1",
        kb_chunking_config={"chunk_size": 32, "chunk_overlap": 4},
    )
    ingested = ingest_document(request)
    assert ingested.chunks
    assert all(row.owner_id == "owner-1" and row.knowledge_base_id == "kb-1" for row in ingested.chunks)
    assert all(row.document_id == "doc-1" and row.source_version == "v1" and not row.index_ready for row in ingested.chunks)
    assert all(row.start >= 0 and row.end >= row.start for row in ingested.chunks)


def test_strategy_chunking_matches_fixed_go_chunker_for_synthetic_golden_cases() -> None:
    cases = json.loads(CHUNK_CASES.read_text(encoding="utf-8"))
    report, oracle = run_chunker_oracle(cases)
    assert report["status"] == "GO_ORACLE_EXECUTED"
    assert report["upstream"] == "3e8b0bfc80b845b2d4b2ed683994748741450a97"
    for case in cases:
        config = case["config"]
        cfg = ChunkConfig(
            chunk_size=config["chunk_size"], chunk_overlap=config["chunk_overlap"],
            separators=tuple(config["separators"]), strategy=config["strategy"],
            token_limit=config.get("token_limit", 0), languages=tuple(config.get("languages", ())),
            enable_parent_child=config.get("parent_child", False),
            parent_chunk_size=config.get("parent_chunk_size", 0), child_chunk_size=config.get("child_chunk_size", 0),
        )
        expected = oracle[case["name"]]
        if config.get("parent_child"):
            actual = split_parent_child(case["text"], cfg)
            assert [_as_chunk(p) for p in actual.parents] == expected["parents"], case["name"]
            assert [{"chunk": _as_chunk(ch.chunk), "parent_index": ch.parent_index} for ch in actual.children] == expected["children"], case["name"]
        else:
            actual, diagnostics = split(case["text"], cfg, with_diagnostics=True)
            assert [_as_chunk(ch) for ch in actual] == expected["chunks"], case["name"]
            assert diagnostics.selected_tier == expected["selected_tier"], case["name"]
            assert list(diagnostics.tier_chain) == expected["tier_chain"], case["name"]
            actual_rejections = [{"tier": tier, "reason": reason} for tier, reason in diagnostics.rejected]
            assert actual_rejections == (expected.get("rejected") or []), case["name"]


def _as_chunk(chunk: Any) -> dict[str, Any]:
    return {"content": chunk.content, "seq": chunk.seq, "start": chunk.start, "end": chunk.end,
            "context_header": chunk.context_header}
