"""S5 offline case freezer. No API, Provider, database or authorization writes.

Builds fixed-context diagnostics, NOT new source documents or RAG results.
Historical bodies/receipts remain immutable. Evaluation expectations never enter
request bodies. The existing Excel parser is reused, not reimplemented.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.weknora_port.ingestion.upstream_excel_parser import ExcelParser

SUMMARY_CONSTRAINT = """

## Identifier-bound facts (isolated S5 candidate)
For an entity associated with an identifier, preserve the source phrase or field
label if the source does not explicitly define its type. Do not infer an entity
type from the identifier, domain knowledge, or a more specific translation.
Apply this to summary, gist, topics and typical_question. Omit missing properties;
do not presuppose them in the question. Original labels may remain untranslated.
"""

SOURCE_HASHES = {
    "north_pier_service_brief.pptx": "382670af50e0e5c7247efaad9ba23dcd0f42db5aec195bbadb74aa766fb62419",
    "site_readings.xlsx": "8791281e99941b5624acf3f40ef26fad81851cbaee3f2391daa9d834a4b7da04",
}
HISTORY = {
    "summary_failure": "7b33910dee444abdb1dc182f15806db7",
    "summary_correct": "6aa2d4aa61d94116a6ee9081b0adb524",
    "table_failure": "3e47e2424d34458e86dd64451b1bd9cd",
    "table_correct": "d16836c522df49fcb14caf1fa3c20049",
}


def sha(value: bytes | str) -> str:
    return hashlib.sha256(value.encode() if isinstance(value, str) else value).hexdigest()


def body_sha(body: dict) -> str:
    return sha(json.dumps(body, ensure_ascii=False, separators=(",", ":"), sort_keys=True, allow_nan=False))


def receipt(path: Path) -> dict:
    """Read saved JSON/SSE; no inferred usage for missing receipts."""
    text = path.read_text(encoding="utf-8")
    events = [json.loads(line[6:]) for line in text.splitlines()
              if line.startswith("data: ") and line[6:] != "[DONE]"] if text.startswith("data:") else [json.loads(text)]
    content, usage, finish, models = [], None, [], set()
    for event in events:
        if event.get("model"):
            models.add(event["model"])
        if event.get("usage"):
            usage = event["usage"]
        for choice in event.get("choices", []):
            message = choice.get("delta", choice.get("message", {}))
            content.append(message.get("content") or "")
            if choice.get("finish_reason"):
                finish.append(choice["finish_reason"])
    return {"response_file": path.name, "response_sha256": sha(path.read_bytes()),
            "models": sorted(models), "answer": "".join(content), "finish_reason": finish,
            "usage": usage if usage is not None else "UNREPORTED",
            "evidence_kind": "HISTORICAL_RESPONSE_NOT_NEW_CALL"}


def capture(directory: Path, stem: str) -> dict:
    path = directory / (stem + ".request.json")
    raw = json.loads(path.read_text(encoding="utf-8"))
    body = raw["body"]
    if body_sha(body) != raw["wire_sha256"]:
        raise ValueError("HISTORICAL_WIRE_SHA_MISMATCH:" + stem)
    if raw["provider"] != "siliconflow" or body["model"] != "Qwen/Qwen3.5-4B":
        raise ValueError("UNEXPECTED_MODEL")
    timing = json.loads((directory / (stem + ".timing.json")).read_text(encoding="utf-8"))
    historical = receipt(directory / (stem + ".response.body"))
    historical["seconds_transport_open_to_read_finished"] = (
        timing["read_finished_monotonic"] - timing["transport_open_started_monotonic"])
    historical["timing_meaning"] = timing["meaning"]
    return {"request_file": path.name, "request_sha256": sha(path.read_bytes()),
            "historical_request_id": raw["request_id"], "model_id": raw["model_id"],
            "owner_sha256": sha(raw["scope"]["owner_id"]),
            "knowledge_base_ids": raw["scope"]["knowledge_base_ids"],
            "document_ids": raw["scope"]["document_ids"],
            "body": body, "wire_sha256": raw["wire_sha256"],
            "receipt": historical}


def replace_once(body: dict, before: str, after: str) -> dict:
    """Only replace one verified source span, never question/system/history."""
    result = deepcopy(body)
    occurrences = sum(m.get("content", "").count(before) for m in result["messages"])
    if not before or occurrences != 1:
        raise ValueError("SOURCE_SPAN_NOT_UNIQUE")
    for message in result["messages"]:
        if before in message.get("content", ""):
            if message["role"] != "user":
                raise ValueError("SOURCE_NOT_USER_DATA")
            message["content"] = message["content"].replace(before, after, 1)
    return result


def literal_excerpt(text: str, max_chars: int = 1024) -> dict:
    """Bounded literal input excerpts, not a summary/profile replacement.

    Skip explicitly model-generated captions, markup-only lines and image paths.
    OCR remains labeled recognition, not certified original PDF characters.
    Every retained line maps exactly to its original input span; no translation.
    """
    if max_chars <= 0:
        raise ValueError("INVALID_EXCERPT_LIMIT")
    excluded = [m.span() for m in re.finditer(r"<image_caption>.*?</image_caption>", text, re.S)]
    spans, used, offset, in_ocr = [], 0, 0, False
    for raw in text.splitlines(keepends=True):
        line = raw.rstrip("\r\n")
        start, end = offset, offset + len(line)
        offset += len(raw)
        if any(start < b and end > a for a, b in excluded):
            continue
        if "<image_ocr>" in line:
            in_ocr = True
        # Subtract markup ranges; retain text offsets, never tag names.
        cursor, regions = 0, []
        for tag in re.finditer(r"<[^>]*>", line):
            regions.append((cursor, tag.start()))
            cursor = tag.end()
        regions.append((cursor, len(line)))
        for left, right in regions:
            value = line[left:right]
            if "<!--" in line or line.lstrip().startswith("![") or not value.strip():
                continue
            a, b = start + left, start + right
            if used + len(value) > max_chars:
                continue
            spans.append({"start": a, "end": b, "text": value,
                          "kind": "ocr_recognition" if in_ocr else "input_text"})
            used += len(value)
        if "</image_ocr>" in line:
            in_ocr = False
    return {"kind": "DERIVED_LITERAL_INPUT_EXCERPT_NOT_PROFILE",
            "input_sha256": sha(text), "limit_chars": max_chars, "spans": spans,
            "all_spans_exact": all(text[s["start"]:s["end"]] == s["text"] for s in spans),
            "not_a_fact_or_language_quality_pass": True}


def parse_excel(data: bytes, semantic_headers: bool):
    parser = ExcelParser(file_name="site_readings.xlsx", file_type="xlsx",
                         xlsx_first_row_as_header=semantic_headers)
    parsed = parser.parse_into_text(data)
    return {"text": parsed.content, "sha256": sha(parsed.content),
            "rows": [{**source, "start": chunk.start, "end": chunk.end}
                     for source, chunk in zip(parser._native_row_sources, parsed.chunks, strict=True)]}


def verified_header_parse(data: bytes, expected_sha: str, *, header_row: int = 1) -> dict:
    """Experiment preconditions, not automatic header detection/production API.

    This narrow fixture path declines merged, formula or ambiguous headers.
    Generated offsets refer ONLY to the new derived text, with original cells
    kept separately. They never claim to be old Chunk or TableRowProof offsets.
    """
    from openpyxl import load_workbook
    if sha(data) != expected_sha:
        raise ValueError("SOURCE_SHA_MISMATCH")
    if type(header_row) is not int or header_row != 1:
        raise ValueError("EXPLICIT_FIRST_HEADER_REQUIRED")
    workbook = load_workbook(BytesIO(data), data_only=False)
    for sheet in workbook:
        headers = [cell.value for cell in sheet[header_row]]
        if sheet.max_row < 2 or any(not isinstance(v, str) or not v.strip() for v in headers):
            raise ValueError("EMPTY_OR_NON_STRING_HEADER")
        normalized = [v.strip().casefold() for v in headers]
        if len(set(normalized)) != len(headers):
            raise ValueError("AMBIGUOUS_DUPLICATE_HEADER")
        if sheet.merged_cells.ranges:
            raise ValueError("MERGED_CELLS_NOT_APPLICABLE")
        if any(cell.data_type == "f" or (isinstance(cell.value, str) and
               re.match(r"^=?(?:_xlfn\.)?(?:DISPIMG|IMAGE)\(", cell.value, re.I))
               for row in sheet for cell in row):
            raise ValueError("FORMULA_NOT_APPLICABLE")
    parsed = parse_excel(data, True)
    proofs = []
    for row in parsed["rows"]:
        sheet = workbook[row["sheet"]]
        row_text = parsed["text"][row["start"]:row["end"]]
        for header in sheet[header_row]:
            cell = sheet.cell(row["row"], header.column)
            if cell.value is None:
                continue
            token = str(header.value).strip() + ": " + str(cell.value)
            if row_text.count(token) != 1:
                raise ValueError("DERIVED_CELL_MAPPING_NOT_EXACT")
            start = row["start"] + row_text.index(token)
            proofs.append({"sheet": sheet.title, "header_cell": header.coordinate,
                           "value_cell": cell.coordinate, "header": header.value,
                           "value": cell.value, "text_start": start,
                           "text_end": start + len(token), "text": token})
    parsed["cell_proofs"] = proofs
    parsed["mapping_kind"] = "DIAGNOSTIC_DERIVED_TEXT_TO_SOURCE_CELLS_NOT_TABLE_ROW_PROOF"
    return parsed


def workbook_projection(data: bytes, columns: tuple[int, ...]) -> tuple[bytes, list[dict]]:
    """Controlled diagnostic projection; original workbook is never saved."""
    from openpyxl import Workbook, load_workbook
    old, new, mapping = load_workbook(BytesIO(data), data_only=True), Workbook(), []
    new.remove(new.active)
    for source in old:
        sheet = new.create_sheet(source.title)
        for row in source:
            sheet.append([row[column - 1].value for column in columns])
            for output_column, source_column in enumerate(columns, 1):
                cell = row[source_column - 1]
                mapping.append({"sheet": source.title, "source_cell": cell.coordinate,
                                "output_row": cell.row, "output_column": output_column,
                                "value": cell.value})
    output = BytesIO()
    new.save(output)
    return output.getvalue(), mapping


def case(identifier, kind, body, candidate, *, history=None, transformation=None, review=None):
    return {"id": identifier, "case_type": kind,
            "baseline": {"body": body, "sha256": body_sha(body),
                         "response": history or "NOT_RUN_NO_HISTORICAL_RESPONSE"},
            "candidate": {"body": candidate, "sha256": body_sha(candidate),
                          "response": "NOT_RUN_BLOCKED_AUTHORIZATION"},
            "transformation": transformation or {"kind": "NONE_ORIGINAL_WIRE"},
            "review_criteria_not_model_input": review,
            "quality_status": "UNVERIFIED_REQUIRES_NEW_AUTHORIZED_RESPONSE"}


def build(source_root: Path) -> dict:
    evidence = source_root / "evidence"
    captures = {name: capture(evidence, stem) for name, stem in HISTORY.items()}
    for name, expected in SOURCE_HASHES.items():
        if sha((source_root / "samples" / name).read_bytes()) != expected:
            raise ValueError("SOURCE_SHA_MISMATCH:" + name)
    original = captures["summary_failure"]["body"]
    source_text = original["messages"][-1]["content"]
    table_bytes = (source_root / "samples/site_readings.xlsx").read_bytes()
    table_old = parse_excel(table_bytes, False)
    table_new = verified_header_parse(table_bytes, SOURCE_HASHES["site_readings.xlsx"])
    if captures["summary_correct"]["body"]["messages"][-1]["content"] != table_old["text"]:
        raise ValueError("EXCEL_NOT_HISTORICAL_TEXT")
    cases = []
    summary_inputs = [("summary_failure", "KNOWN_FAILURE", original, None),
                      ("summary_correct", "KNOWN_CORRECT_CONTROL", captures["summary_correct"]["body"], None)]
    # Delete identifier-bearing whole lines in each source layer. This is an
    # explicit adversarial fixture mutation, never a production recognition rule.
    omitted = []
    keep = []
    for match in re.finditer(r".*(?:\n|$)", source_text):
        if not match.group():
            continue
        if "NP-204" in match.group():
            omitted.append({"start": match.start(), "end": match.end(), "text": match.group()})
        else:
            keep.append(match.group())
    summary_inputs.append(("summary_missing_identity", "CONSTRUCTED_MISSING_EVIDENCE",
        replace_once(original, source_text, "".join(keep)), {"kind": "DELETE_DIAGNOSTIC_IDENTIFIER_LINES", "removed": omitted}))
    mixed = "<source_a>\n" + source_text + "\n</source_a>\n<source_b>\n" + table_old["text"] + "</source_b>"
    summary_inputs.append(("summary_distractor", "CONSTRUCTED_MULTI_RECORD_DISTRACTOR",
        replace_once(original, source_text, mixed), {"kind": "CONCAT_TWO_REAL_INPUTS_SEPARATE_BOUNDARIES",
        "source_a_sha256": sha(source_text), "source_b_sha256": table_old["sha256"]}))
    for identifier, kind, body, transformation in summary_inputs:
        candidate = deepcopy(body)
        candidate["messages"][0]["content"] += SUMMARY_CONSTRAINT
        record = case(identifier, kind, body, candidate,
            history=captures[identifier]["receipt"] if identifier in captures else None,
            transformation=transformation,
            review="No unsupported subtype or recovered missing identity in any profile field; retain actual identifiers, relations, dates and configured language.")
        record["literal_alternative"] = literal_excerpt(body["messages"][-1]["content"])
        cases.append(record)
    for identifier in ("table_failure", "table_correct"):
        old = captures[identifier]["body"]
        cases.append(case(identifier, "KNOWN_FAILURE" if identifier.endswith("failure") else "KNOWN_CORRECT_CONTROL",
            old, replace_once(old, table_old["text"], table_new["text"]), history=captures[identifier]["receipt"],
            transformation={"kind": "EXISTING_EXPLICIT_HEADER_OPTION", "baseline": table_old, "candidate": table_new},
            review="Keep Site and Unit distinct; retain matching row, reading and CheckedOn, no new device type or RecordedOn claim."))
    for identifier, kind, columns in (
        ("table_missing_unit", "CONSTRUCTED_MISSING_EVIDENCE", (1, 3, 4)),
        ("table_column_distractor", "CONSTRUCTED_REORDERED_COLUMNS", (2, 1, 4, 3)),
    ):
        data, mapping = workbook_projection(table_bytes, columns)
        old, new = parse_excel(data, False), verified_header_parse(data, sha(data))
        baseline = replace_once(captures["table_failure"]["body"], table_old["text"], old["text"])
        cases.append(case(identifier, kind, baseline, replace_once(baseline, old["text"], new["text"]),
            transformation={"kind": "IN_MEMORY_SOURCE_CELL_PROJECTION_NOT_NEW_DOCUMENT", "source_columns": columns,
                            "cell_mapping": mapping, "baseline": old, "candidate": new},
            review="Missing Unit must remain unknown, not answered with Site. Reordered columns must bind by actual header, never position assumption."))
    return {"schema_version": "s5-fixed-facts-v1", "scope": "OFFLINE_FIXED_CONTEXT_DIAGNOSTICS_NOT_RAG_ACCEPTANCE",
            "new_provider_calls": 0, "new_tokens": 0, "new_unknown": 0,
            "authorization": "CLOSED_NOT_REOPENED", "production_changed": False,
            "source_files_sha256": SOURCE_HASHES,
            "historical_inputs": {name: {k: v for k, v in item.items() if k not in ("body", "receipt")}
                                  for name, item in captures.items()},
            "cases": cases, "requested_future_new_calls": 12,
            "future_call_accounting": "8 candidate generations + 4 baselines for constructed cases; 4 saved historical controls reused non-contemporaneously, not resent.",
            "not_proven": ["New model quality", "Normal API retrieval behavior", "Production default suitability"]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = build(args.source_root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "OFFLINE_CASE_PACK_BUILT_NOT_QUALITY_PASS", "cases": len(result["cases"]),
                      "provider_calls": 0, "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
