"""Create small synthetic S3 parser samples and run the local ParserRegistry.

These are acceptance fixtures, not official benchmark Gold. All output is
written outside the repository under E:\\WK\\run\\s3-product-20261019\\samples.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = Path(r"E:\WK\run\s3-product-20261019")
OUT = RUN_DIR / "samples"
CONFIG = Path(r"E:\WK\run\p81-retirement-20261018\local.json")
LOCAL_PYTHON = Path(r"E:\RAG quention\.venv\Scripts\python.exe")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def create_scan_card(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 850), "#fbfaf5")
    draw = ImageDraw.Draw(image)
    # Use a common Windows sans font if present; Pillow's built-in font is a
    # deterministic fallback. The short uppercase labels are OCR-friendly.
    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    font = ImageFont.truetype(str(font_path), 66) if font_path.is_file() else ImageFont.load_default()
    draw.rounded_rectangle((55, 55, 1145, 795), radius=28, outline="#25364a", width=8, fill="white")
    draw.text((115, 145), "FIELD CHECK CARD", font=font, fill="#17324d")
    draw.text((115, 310), "SITE: NORTH PIER", font=font, fill="#111111")
    draw.text((115, 435), "UNIT: NP-204", font=font, fill="#111111")
    draw.text((115, 560), "READING: 37", font=font, fill="#111111")
    draw.line((115, 680, 1080, 680), fill="#777777", width=5)
    image.save(path, format="PNG", optimize=True)


def build_pdf(path: Path, scanned_image: Path) -> None:
    import fitz

    with fitz.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_text((60, 70), "North Pier Maintenance Note", fontsize=19)
        page.insert_text((60, 112), "Inspection date: 2026-10-12", fontsize=12)
        page.insert_text((60, 150), "The scanned field check card is filed on page 2.", fontsize=12)
        page.insert_text((60, 205), "Item | Count", fontsize=12)
        page.insert_text((60, 225), "Valves | 4", fontsize=12)
        page.insert_text((60, 245), "Meters | 2", fontsize=12)
        scan_page = pdf.new_page(width=612, height=792)
        scan_page.insert_image(scan_page.rect, filename=str(scanned_image))
        pdf.save(path, garbage=4, deflate=True)


def build_pptx(path: Path, image_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.5), Inches(0.7))
    title.text_frame.text = "North Pier Service Brief"
    title.text_frame.paragraphs[0].font.size = Pt(24)
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(8.8), Inches(1.2))
    body.text_frame.text = "The maintenance team checked unit NP-204."
    slide.shapes.add_picture(str(image_path), Inches(1.0), Inches(3.2), width=Inches(5.0))

    slide2 = deck.slides.add_slide(deck.slide_layouts[6])
    title2 = slide2.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.5), Inches(0.7))
    title2.text_frame.text = "Visit Summary"
    title2.text_frame.paragraphs[0].font.size = Pt(24)
    body2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(8.8), Inches(1.5))
    body2.text_frame.text = "The next review is scheduled for 2026-10-20."
    deck.save(path)


def build_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "SiteReadings"
    rows = [
        ("Site", "Unit", "Reading", "CheckedOn"),
        ("North Pier", "NP-204", 37, "2026-10-12"),
        ("East Gate", "EG-118", 24, "2026-10-12"),
        ("West Yard", "WY-309", 51, "2026-10-13"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
    sheet.freeze_panes = "A2"
    book.save(path)


def main() -> int:
    if not CONFIG.is_file() or not LOCAL_PYTHON.is_file():
        raise RuntimeError("REQUIRED_LOCAL_CONFIG_OR_PYTHON_MISSING")
    sys.path.insert(0, str(ROOT))
    from scripts.local_runtime import environment, load_config

    cfg = load_config(CONFIG)
    # Reuse only the configured import paths; do not print the runtime config.
    runtime_env = environment(cfg)
    configured = runtime_env.get("PYTHONPATH", "")
    for entry in reversed([part for part in configured.split(os.pathsep) if part]):
        sys.path.insert(0, entry)

    OUT.mkdir(parents=True, exist_ok=True)
    if (OUT / "FROZEN").exists():
        raise RuntimeError("SAMPLE_ALREADY_FROZEN_NO_REGENERATION")
    scan = OUT / "north_pier_scan_card.png"
    pdf = OUT / "north_pier_inspection.pdf"
    pptx = OUT / "north_pier_service_brief.pptx"
    xlsx = OUT / "site_readings.xlsx"
    create_scan_card(scan)
    build_pdf(pdf, scan)
    build_pptx(pptx, scan)
    build_xlsx(xlsx)

    files = [
        {"id": "mixed_pdf", "path": pdf.name, "sha256": sha256(pdf), "format": "pdf",
         "pages": [{"page": 1, "kind": "native_text_and_small_table"},
                   {"page": 2, "kind": "full_page_scanned_image"}],
         "expected_assets": [{"identity": "north_pier_inspection.pdf:page-2:scanned-image",
                              "page": 2, "source_type": "scanned_pdf"}]},
        {"id": "two_slide_pptx", "path": pptx.name, "sha256": sha256(pptx), "format": "pptx",
         "slides": [{"slide": 1, "text": "North Pier Service Brief; unit NP-204",
                     "embedded_image": "north_pier_scan_card.png"},
                    {"slide": 2, "text": "Visit Summary; 2026-10-20"}],
         "expected_assets": [{"identity": "north_pier_service_brief.pptx:slide-1:embedded-image-1",
                              "slide": 1, "source_type": "pptx_media"}]},
        {"id": "site_readings", "path": xlsx.name, "sha256": sha256(xlsx), "format": "xlsx",
         "sheets": [{"name": "SiteReadings", "header_row": 1, "data_rows": 3}],
         "expected_assets": []},
    ]
    save_json(OUT / "source_manifest.json", {
        "classification": "synthetic acceptance fixture; not official Gold",
        "facts": {
            "pdf": {"title": "North Pier Maintenance Note", "inspection_date": "2026-10-12",
                    "table": {"Valves": 4, "Meters": 2}, "scanned_card": {
                        "site": "NORTH PIER", "unit": "NP-204", "reading": 37}},
            "pptx": {"title": "North Pier Service Brief", "unit": "NP-204",
                     "slide_2_review_date": "2026-10-20"},
            "xlsx": {"sheet": "SiteReadings", "rows": [
                {"Site": "North Pier", "Unit": "NP-204", "Reading": 37, "CheckedOn": "2026-10-12"},
                {"Site": "East Gate", "Unit": "EG-118", "Reading": 24, "CheckedOn": "2026-10-12"},
                {"Site": "West Yard", "Unit": "WY-309", "Reading": 51, "CheckedOn": "2026-10-13"},
            ]},
            "image_file": {"path": scan.name, "sha256": sha256(scan),
                           "embedded_image_identity": "north_pier_service_brief.pptx:slide-1:embedded-image-1"},
        }, "files": files,
    })
    save_json(OUT / "questions_expected.json", {
        "classification": "synthetic acceptance questions; expected facts for post-parse comparison",
        "questions": [
            {"id": "s3-q1", "file_id": "mixed_pdf", "question": "What was the inspection date?",
             "expected": "2026-10-12", "source_location": {"page": 1}},
            {"id": "s3-q2", "file_id": "mixed_pdf", "question": "What reading is shown on the scanned card?",
             "expected": "37", "source_location": {"page": 2, "image_identity": "north_pier_inspection.pdf:page-2:scanned-image"}},
            {"id": "s3-q3", "file_id": "two_slide_pptx", "question": "What site and reading are visible on the card in the service brief?",
             "expected": {"site": "North Pier", "reading": 37}, "source_location": {"slide": 1}},
            {"id": "s3-q4", "file_id": "two_slide_pptx", "question": "When is the next review scheduled?",
             "expected": "2026-10-20", "source_location": {"slide": 2}},
            {"id": "s3-q5", "file_id": "site_readings", "question": "What is the reading for East Gate?",
             "expected": 24, "source_location": {"sheet": "SiteReadings", "row": 3}},
            {"id": "s3-q6", "file_id": "site_readings", "question": "Which unit has a reading of 51?",
             "expected": "WY-309", "source_location": {"sheet": "SiteReadings", "row": 4}},
        ],
    })

    from backend.weknora_port.ingestion.core import DocumentRequest, ParserRegistry

    registry = ParserRegistry()
    parsed = []
    try:
        for item in files:
            path = OUT / item["path"]
            # Empty parser_engine means normal extension-default dispatch.
            engine, document = registry.parse(DocumentRequest(filename=path.name, data=path.read_bytes()))
            if document.failures:
                failure = document.failures[0]
                raise RuntimeError(f"{item['id']}:{engine}:{failure.code}:{failure.exception_type}:{failure.message}")
            parsed.append({"file_id": item["id"], "parser": engine,
                           "text_characters": len(document.content), "asset_count": len(document.assets),
                           "asset_identities": [
                               {"path": asset.path, "source_type": asset.source_type,
                                "page_number": asset.page_number, "media_type": asset.media_type}
                               for asset in document.assets],
                           "text_preview": document.content[:1200]})
    except Exception as exc:
        save_json(OUT / "first_error.json", {"type": type(exc).__name__, "message": str(exc)})
        raise
    save_json(OUT / "parse_results.json", {"mode": "local ParserRegistry; no OCR/provider/service/database",
                                             "documents": parsed})
    print(json.dumps({"output": str(OUT), "documents": len(files), "questions": 6,
                      "image_assets": sum(len(row["asset_identities"]) for row in parsed),
                      "parsers": [row["parser"] for row in parsed]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
